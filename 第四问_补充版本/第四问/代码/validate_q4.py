"""Independent Q4 replay: direct timestamp comparisons, exhaustive label products.

Does not import the partition solver, union-find, heap coloring or rank functions.
Uses bundled openpyxl only to read authoritative input workbooks.
"""
from collections import Counter, defaultdict
import argparse
import csv
from decimal import Decimal, ROUND_CEILING
from fractions import Fraction
import hashlib
import itertools
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '第四问/结果'
KEYS = ('U_A', 'U_B', 'U_C', 'B_A', 'B_B', 'B_C', 'R', 'RE')
AIR = ('U_A', 'U_B', 'U_C', 'R')
POWER = ('B_A', 'B_B', 'B_C', 'RE')


def digest(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


class Audit:
    def __init__(self):
        self.count = 0
        self.errors = []

    def check(self, ok, message):
        self.count += 1
        if not ok:
            self.errors.append(message)


def main(output=OUT):
    OUT = Path(output).resolve()
    a = Audit()
    result = json.loads((OUT / 'solution_q4.json').read_text('utf8'))
    original = Path(result['source_solution'])
    frozen = OUT.parent / '数据/q3_frozen_solution.json'
    a.check(digest(original) == digest(frozen) == result['source_solution_sha256'], 'Q3 source bytes changed')
    q3 = json.loads(frozen.read_text('utf8'))
    old_validation = json.loads((OUT / 'q3_baseline_validation.json').read_text('utf8'))
    a.check(old_validation['passed'] and old_validation['solution_sha256'] == digest(frozen), 'Q3 proof binding')
    a.check(digest(OUT / 'q3_baseline_validation.json') == result['source_validation_sha256'], 'Q3 proof bytes changed')
    for p, h in result['source_input_hashes'].items():
        a.check(digest(ROOT / p.replace('\\', '/')) == h, f'Input hash {p}')
    base = ROOT / '数据/无人机应急物资运输基础数据'
    w = openpyxl.load_workbook(base / '运输无人机数据.xlsx', data_only=True)
    rows = list(w['数据'].values)
    inv = {f'U_{t}': sum(r[1] == t for r in rows[8:16]) for t in 'ABC'}
    inv.update({f'B_{r[0]}': int(r[1]) for r in rows[19:22]})
    capacities = {r[0]: float(r[8]) for r in rows[2:5]}
    charges = {r[0]: float(r[2]) for r in rows[19:22]}
    w.close()
    w = openpyxl.load_workbook(base / '中继无人机数据.xlsx', data_only=True)
    rrows = list(w['数据'].values)
    inv['R'], inv['RE'] = len(rrows[6:8]), int(rrows[11][1])
    capacities['R'], charges['R'] = rrows[2][7], rrows[11][2]
    turnaround = rrows[2][11]
    w.close()
    w = openpyxl.load_workbook(base / '物资需求与配送时限.xlsx', data_only=True)
    boxes = {r[0]: r[1] for r in list(w['逐箱货箱清单'].values)[1:] if r[0]}
    w.close()
    a.check(inv == result['inventory'], 'Source inventory mismatch')
    a.check(Counter(b for t in q3['transport'] for b in t['boxes']) == Counter({b: 1 for b in boxes}), '80 original boxes')
    nodes = sorted(set(boxes.values()))
    transport = {t['id']: t for t in q3['transport']}
    relays = {r['id']: r for r in q3['relay']}
    edges = {r: set() for r in relays}
    assoc_count = 0
    for t in q3['transport']:
        for c in t['communication']:
            if not c['relay']:
                continue
            st = c.get('station', t['station'])
            lo = t['start_s'] + c['start']
            hi = t['start_s'] + c['end']
            active = [r for r in q3['relay'] if r['station'] == st
                      and r['service_start_s'] - 1e-6 <= lo <= hi <= r['service_end_s'] + 1e-6]
            a.check(len(active) == 1, f'Active communication binding {t["id"]}')
            if len(active) == 1:
                edges[active[0]['id']].add(t['id'])
            assoc_count += 1
    a.check({k: sorted(v) for k, v in edges.items()} == result['relay_transport_edges'], 'Relay edges mismatch')
    a.check(len(result['communication_associations']) == assoc_count, 'Communication segment count')

    def charge_ms(e, kind):
        # Independent SOC-based two-phase charge integration.
        soc = Decimal(1) - Decimal(str(e)) / Decimal(str(capacities[kind]))
        fast = max(Decimal('.9') - soc, Decimal(0)) / Decimal('.9') * Decimal('.65')
        slow = (Decimal(1) - max(soc, Decimal('.9'))) / Decimal('.1') * Decimal('.35')
        return int((Decimal(str(charges[kind])) * (fast + slow) * 1000).to_integral_value(rounding=ROUND_CEILING))

    intervals = {key: [] for key in KEYS}
    for t in q3['transport']:
        start = round(t['start_s'] * 1000)
        intervals['U_' + t['model']].append((t['id'], start, round(t['resource_end_s']*1000), t['drone']))
        ready = round(t['return_s']*1000) + charge_ms(t['energy_kwh'], t['model'])
        a.check(round(t['battery_ready_s']*1000) >= ready, f'Transport charge {t["id"]}')
        intervals['B_' + t['model']].append((t['id'], start, round(t['battery_ready_s']*1000), t['battery']))
    for r in q3['relay']:
        start = round(r['start_s']*1000)
        end = round(r['resource_end_s']*1000)
        a.check(end >= round((r['return_s']+turnaround)*1000), f'Relay turnaround {r["id"]}')
        ready = round(r['return_s']*1000) + charge_ms(r['energy_kwh'], 'R')
        ready = max([ready]+[round(r[k]*1000) for k in ('component_ready_s','battery_ready_s','charge_end_s') if k in r])
        intervals['R'].append((r['id'],start,end,r['drone']))
        intervals['RE'].append((r['id'],start,ready,r['component']))
    for key in KEYS:
        exported = [(x['task'],x['start_ms'],x['end_ms'],x['original_id']) for x in result['resource_intervals'][key]]
        a.check(sorted(intervals[key]) == sorted(exported), f'Interval reconstruction {key}')

    # Graph reachability separately reconstructs both block systems.
    def blocks(strict):
        adjacent = {s: {s} for s in nodes}
        lists = [t['order'] for t in q3['transport']]
        if strict:
            lists += [sorted({s for t in ts for s in transport[t]['order']}) for ts in edges.values()]
        for ss in lists:
            for s in ss:
                adjacent[s].update(ss)
        left, out = set(nodes), []
        while left:
            stack, reached = [min(left)], set()
            while stack:
                s = stack.pop()
                if s not in reached:
                    reached.add(s)
                    stack.extend(adjacent[s] - reached)
            left -= reached
            out.append(sorted(reached))
        return sorted(out)

    a.check(blocks(False) == result['transport_blocks'], 'Transport connected components')
    a.check(blocks(True) == result['strict_blocks'], 'Joint connected components')
    cache = {}

    def group_counts(services):
        key = tuple(sorted(services))
        if key in cache:
            return cache[key]
        ts = {t for t, row in transport.items() if row['order'][0] in services}
        rs = {r for r, targets in edges.items() if targets & ts}
        counts, ids = {}, {}
        for resource in KEYS:
            ints = [i for i in intervals[resource] if i[0] in ts | rs]
            counts[resource] = max((sum(start <= t < end for _,start,end,_ in ints) for _,t,_,_ in ints), default=0)
            ids[resource] = len({i[3] for i in ints})
        work = sum(round((transport[t]['return_s']-transport[t]['start_s'])*1000) for t in ts)
        cache[key] = counts, ids, work, ts, rs
        return cache[key]

    rows = json.loads((OUT / 'all_partitions.json').read_text('utf8'))
    grouped_rows = defaultdict(list)
    for row in rows:
        policy, k, identity = row['policy'], row['K'], row['identity']
        grouped_rows[policy,k,identity].append(row)
        a.check(sorted(s for g in row['groups'] for s in g) == nodes and all(row['groups']), f'Partition coverage {row["partition_id"]}')
        groups = [group_counts(g) for g in row['groups']]
        for g, (_, _, _, ts, rs) in zip(row['groups'], groups):
            a.check(all(set(transport[t]['order']) <= set(g) for t in ts), 'Transport split')
        if policy == 'strict':
            a.check(all(n == 1 for n in Counter(r for g in groups for r in g[4]).values()), 'Relay task duplicated in primary')
        total = {r: sum(g[0 if identity == 'minimum' else 1][r] for g in groups) for r in KEYS}
        gap = {r: max(total[r]-inv[r],0) for r in KEYS}
        work = [g[2] for g in groups]
        var = Fraction(k*sum(x*x for x in work),sum(work)**2)-1
        a.check(total == row['total'] and gap == row['gap'], f'Peak/shortage {row["partition_id"]}')
        a.check(Fraction(row['workload_cv_squared_fraction']) == var, f'Workload {row["partition_id"]}')
        energy = sum(sum(transport[t]['energy_kwh'] for t in g[3]) + sum(relays[r]['energy_kwh'] for r in g[4]) for g in groups)
        a.check(abs(energy-row['total_energy_kwh']) < 1e-9, 'Relay duplicate energy accounting')

    def canonical(groups):
        return tuple(sorted(tuple(sorted(g)) for g in groups))

    for (policy,k,identity), candidates in grouped_rows.items():
        atomic = blocks(policy == 'strict')
        exhaustive = set()
        for assignment in itertools.product(range(k), repeat=len(atomic)):
            if len(set(assignment)) != k:
                continue
            gs = [[s for j,block in enumerate(atomic) if assignment[j] == g for s in block] for g in range(k)]
            exhaustive.add(canonical(gs))
        a.check(len(candidates) == len(exhaustive) and {canonical(r['groups']) for r in candidates} == exhaustive, f'Enumeration complete {policy}/{k}/{identity}')

        def score(row):
            return (sum(row['gap'][r] for r in AIR),sum(row['gap'][r] for r in POWER),
                    sum(row['total'][r] for r in AIR),sum(row['total'][r] for r in POWER),
                    Fraction(row['workload_cv_squared_fraction']))

        rec = result['recommendations'][f'{policy}_{k}_{identity}']
        a.check(score(rec) == min(map(score,candidates)), f'Exact optimum {policy}/{k}/{identity}')
        a.check(rec['groups'] == next(r['groups'] for r in candidates if r['partition_id'] == rec['partition_id']), 'Recommendation membership')
        vectors = [(tuple(row['total'][r] for r in KEYS) + (Fraction(row['workload_cv_squared_fraction']),), row['partition_id']) for row in candidates]
        frontier = {pid for v,pid in vectors if not any(all(x<=y for x,y in zip(u,v)) and u!=v for u,_ in vectors)}
        a.check(frontier == set(rec['pareto_partition_ids']), f'Independent Pareto frontier {policy}/{k}/{identity}')
        for g in rec['group_details']:
            counts, ids, _, ts, rs = group_counts(g['services'])
            a.check(counts == g['resource_minimum'] and ids == g['original_id_count'], 'Recommendation group counts')
            for resource, assigns in g['assignments'].items():
                a.check(Counter(x['task'] for x in assigns) == Counter(i[0] for i in intervals[resource] if i[0] in ts | rs), 'No missing resource assignment')
                a.check(len({x['slot'] for x in assigns}) == counts[resource], 'Coloring attains peak')
                for x,y in itertools.combinations(assigns,2):
                    a.check(x['slot'] != y['slot'] or x['end_ms'] <= y['start_ms'] or y['end_ms'] <= x['start_ms'], 'Allocated resource overlap')
                witness = g['peak_witnesses'][resource]
                if witness:
                    tick = round(witness['time_s']*1000)
                    actual = {i[0] for i in intervals[resource] if i[0] in ts|rs and i[1] <= tick < i[2]}
                    a.check(actual == set(witness['tasks']) and len(actual) == counts[resource], 'Peak lower-bound witness')

    with (OUT/'resource_assignments.csv').open(encoding='utf-8-sig',newline='') as f:
        allocation=list(csv.DictReader(f))
    for k in (2,3):
        selected=[r for r in allocation if int(r['K'])==k]
        physical=defaultdict(list)
        for r in selected:
            physical[r['resource_type'],r['local_resource']].append(r)
        for (typ,name), uses in physical.items():
            a.check(len({r['group'] for r in uses})==1, f'Cross-group resource use {k}/{name}')
            uses.sort(key=lambda x:float(x['start_s']))
            for x,y in zip(uses,uses[1:]):
                a.check(float(x['end_s'])<=float(y['start_s']), f'Physical assignment overlap {name}')
        rec=result['recommendations'][f'strict_{k}_minimum']
        a.check({r:sum(typ==r and name.startswith('ADD-') for typ,name in physical) for r in KEYS}==rec['gap'],'Physical additional stock matches gap')
        a.check({r:sum(typ==r for typ,name in physical) for r in KEYS}==rec['total'],'Physical configuration totals')
        for i,g in enumerate(rec['group_details'],1):
            _,_,_,ts,rs=group_counts(g['services'])
            for resource in KEYS:
                got=[r for r in selected if r['group']==f'G{i}' and r['resource_type']==resource]
                expected=[row for row in intervals[resource] if row[0] in ts|rs]
                a.check(Counter(r['task'] for r in got)==Counter(row[0] for row in expected),'Physical allocation covers exact tasks')
                a.check(len({r['local_resource'] for r in got})==g['resource_minimum'][resource],'Group physical stock count')
                lookup={row[0]:row for row in expected}
                mapping=defaultdict(set)
                for r in got:
                    row=lookup[r['task']]
                    a.check(round(float(r['start_s'])*1000)==row[1] and round(float(r['end_s'])*1000)==row[2] and r['q3_original_id']==row[3],'Physical allocation preserves times and original identity')
                    mapping[r['q3_original_id']].add(r['local_resource'])
                a.check(all(len(v)==1 for v in mapping.values()),'Primary preserves within-group Q3 identity reuse')

    # Boundary cases distinguish the intended half-open rule and charging.
    edge=[(0,1000),(1000,2000)]
    a.check(max(sum(s<=t<e for s,e in edge) for t in (0,1000))==1,'Endpoint reuse rule')
    extended=[(0,1500),(1000,2500)]
    a.check(max(sum(s<=t<e for s,e in extended) for t in (0,1000))==2,'Charging must prevent early reuse')
    report={'passed':not a.errors,'check_count':a.count,'error_count':len(a.errors),'errors':a.errors,
            'solution_sha256':digest(OUT/'solution_q4.json'),'validator_sha256':digest(__file__),
            'all_partition_rows':len(rows),'q3_sha256':digest(frozen),
            'scope':'Independent partition/resource replay; continuous physical feasibility inherited from current hash-bound Q3 validation.'}
    (OUT/'validation_q4.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    if a.errors:
        raise SystemExit(1)


if __name__=='__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results', type=Path, default=OUT)
    main(parser.parse_args().results)
