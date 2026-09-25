"""Read back the exported XLSX; never save or alter source workbooks."""
import csv
import argparse
import hashlib
import json
from pathlib import Path
import re
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--results',type=Path,default=ROOT/'第四问/结果')
parser.add_argument('--output',type=Path,default=ROOT/'outputs/question4_partition_20260925')
args=parser.parse_args()
RESULT=args.results.resolve()
OUTPUT=args.output.resolve()
BOOK=OUTPUT/'第四问_任务分区与资源配置.xlsx'
data=json.loads((RESULT/'solution_q4.json').read_text('utf8'))
q3=json.loads((RESULT.parent/'数据/q3_frozen_solution.json').read_text('utf8'))
allrows=json.loads((RESULT/'all_partitions.json').read_text('utf8'))
v=openpyxl.load_workbook(BOOK,data_only=True)
f=openpyxl.load_workbook(BOOK,data_only=False)
errors=[]
checks=0


def check(ok,why):
    global checks
    checks+=1
    if not ok:
        errors.append(why)


def eq(got,wanted,why):
    check(isinstance(got,(int,float)) and abs(got-wanted)<1e-7 if isinstance(wanted,(int,float)) else got==wanted,f'{why}: {got!r} / {wanted!r}')


keys=data['resource_keys']
check(v.sheetnames==['结果比较','Q4_分区配置','任务明细','资源占用','分区备选','库存与口径'],'Sheet coverage/order')
template=openpyxl.load_workbook(ROOT/'结果提交模板.xlsx',data_only=True)
eq([v['Q4_分区配置'].cell(1,c).value for c in range(1,12)],[template['Q4_分区配置'].cell(1,c).value for c in range(1,12)],'Official Q4 headers')
template.close()
row=2
for k in (2,3):
    rec=data['recommendations'][f'strict_{k}_minimum']
    for i,g in enumerate(rec['group_details'],1):
        s=v['Q4_分区配置']
        eq(s.cell(row,1).value,k,'K');eq(s.cell(row,2).value,f'G{i}','Group')
        eq(re.findall(r'S\d{3}',s.cell(row,3).value),g['services'],'Service IDs survive wrapping')
        for j,r in enumerate(keys,4):
            eq(s.cell(row,j).value,g['resource_minimum'][r],f'Configuration {k}/{i}/{r}')
        for j,wanted in enumerate([len(g['transport']),g['box_count'],g['transport_work_ms']/1000,g['relay_work_ms']/1000,g['population'],g['mass_kg']],12):
            eq(s.cell(row,j).value,wanted,f'Group support {row}/{j}')
        row+=1
for i,r in enumerate(keys,5):
    two=data['recommendations']['strict_2_minimum'];three=data['recommendations']['strict_3_minimum']
    wanted=[data['inventory'][r],data['global_group']['resource_minimum'][r],two['total'][r],two['redundancy_vs_global_minimum'][r],two['gap'][r],three['total'][r],three['redundancy_vs_global_minimum'][r],three['gap'][r]]
    for j,x in enumerate(wanted,3):
        eq(v['结果比较'].cell(i,j).value,x,f'Summary {i}/{j}')
for k,c in ((2,2),(3,3)):
    eq(v['结果比较'].cell(20,c).value,data['recommendations'][f'strict_{k}_minimum']['workload_cv'],'CV')
    eq(v['结果比较'].cell(21,c).value,q3['metrics']['joint_finish_s'],'Makespan')
    eq(v['结果比较'].cell(22,c).value,q3['metrics']['total_energy_kwh'],'Energy')
for i,t in enumerate(q3['transport']+q3['relay'],5):
    relay=i>=5+len(q3['transport'])
    actual=[v['任务明细'].cell(i,c).value for c in range(1,13)]
    groups=[]
    for k in (2,3):
        rec=data['recommendations'][f'strict_{k}_minimum']
        groups.append(next(f'G{j}' for j,g in enumerate(rec['group_details'],1) if t['id'] in g['relay' if relay else 'transport']))
    expected=['中继' if relay else '运输',t['id'],
              '、'.join(data['relay_transport_edges'][t['id']]) if relay else ' → '.join(t['order']),
              'R' if relay else t['model'],*groups,0 if relay else len(t['boxes']),t['start_s'],t['return_s'],
              round(t['return_s']-t['start_s'],3),t['energy_kwh'],
              f'站{t["station"]}' if relay else '、'.join(r for r,ts in data['relay_transport_edges'].items() if t['id'] in ts) or '直连']
    for j,(x,y) in enumerate(zip(actual,expected),1):eq(x,y,f'Task {i}/{j}')
with (RESULT/'resource_assignments.csv').open(encoding='utf-8-sig',newline='') as src:
    resources=list(csv.DictReader(src))
for i,r in enumerate(resources,5):
    expected=[int(r['K']),r['group'],data['resource_labels'][r['resource_type']],r['local_resource'],r['q3_original_id'],r['task'],
              '现有' if r['supply']=='existing' else '待补',float(r['start_s']),float(r['end_s']),round(float(r['end_s'])-float(r['start_s']),3)]
    for j,wanted in enumerate(expected,1):eq(v['资源占用'].cell(i,j).value,wanted,f'Resource {i}/{j}')
check(v['资源占用'].cell(5+len(resources),1).value is None,'No extra resource row')
alts=[r for r in allrows if r['policy']=='strict' and r['identity']=='minimum']
for k in (2,3):
    alts.extend([data['recommendations'][f'clone_{k}_minimum'],data['recommendations'][f'clone_{k}_minimum']['balance_first']])
for i,r in enumerate(alts,5):
    eq(v['分区备选'].cell(i,2).value,r['K'],'Alternative K')
    eq(re.findall(r'S\d{3}',v['分区备选'].cell(i,4).value),[s for g in r['groups'] for s in g],'Alternative groups')
    for j,wanted in enumerate([r['total'][k] for k in keys]+[r['workload_cv'],r['relay_sorties_executed'],r['total_energy_kwh']],5):
        eq(v['分区备选'].cell(i,j).value,wanted,f'Alternative {i}/{j}')
formula_count=0
for s in f:
    check(not s.sheet_view.showGridLines,f'Gridlines {s.title}')
    for row in s:
        for c in row:
            cached=v[s.title][c.coordinate]
            check(cached.data_type!='e',f'Excel error {s.title}/{c.coordinate}')
            if c.data_type=='f':
                formula_count+=1
                check(isinstance(cached.value,(int,float)),f'Formula cache {s.title}/{c.coordinate}')
                check(not any(term in c.value for term in ('#REF!','[1]','http://')),f'Broken/external formula {c.coordinate}')
for name in ('任务明细','资源占用'):
    check(f[name].freeze_panes=='A5',f'Frozen headings {name}')
for name in v.sheetnames:
    check(len(f[name].tables)==1,f'Filter table {name}')
report={'passed':not errors,'check_count':checks,'error_count':len(errors),'errors':errors,'formulas':formula_count,
        'resource_interval_rows':len(resources),'task_rows':len(q3['transport'])+len(q3['relay']),'config_rows':sum(data['recommendations'][f'strict_{k}_minimum']['K'] for k in (2,3)),
        'output_sha256':hashlib.sha256(BOOK.read_bytes()).hexdigest(),
        'source_solution_sha256':hashlib.sha256((RESULT/'solution_q4.json').read_bytes()).hexdigest(),
        'engine':'artifact-tool recalculation and openpyxl saved-cache readback; native Excel/WPS not launched'}
(OUTPUT/'workbook_readback_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
if errors:raise SystemExit(1)
