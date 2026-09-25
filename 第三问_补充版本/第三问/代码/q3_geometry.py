"""Independent Q3 terrain, bidirectional radio budgets and continuous certificates.

Reads original attachments only. Q1 supplies two generic geometry functions;
no previous solution files are consumed. Heights are absolute metres.
"""
from __future__ import annotations

import argparse
import json
import math
from functools import lru_cache
from pathlib import Path

import numpy as np
import openpyxl
from PIL import Image
from solve_q1_batching import cells_on_segment, local_plane

ROOT = Path(__file__).resolve().parents[1]


class Geometry:
    def __init__(self, root=ROOT):
        self.root = Path(root)
        base = self.root / '数据' / '无人机应急物资运输基础数据'
        wb = openpyxl.load_workbook(base/'调度中心与服务区.xlsx', read_only=True, data_only=True)
        raw = {}
        for row in wb['数据'].values:
            if isinstance(row[0], str) and (row[0] == 'O01' or row[0].startswith('S0')):
                raw[row[0]] = dict(id=row[0], name=row[1], lon=float(row[2]), lat=float(row[3]), elevation_m=float(row[4]))
        wb.close()
        self.coordinates, self.origin = local_plane(raw)
        self.nodes = {k: dict(**v, x=self.coordinates[k]['x_m'], y=self.coordinates[k]['y_m'], z=v['elevation_m']) for k,v in raw.items()}
        self.dem_path = next((self.root/'数据').rglob('*.tif'))
        image = Image.open(self.dem_path)
        self.dem = np.asarray(image, dtype=float)
        scale, tie, keys = image.tag_v2[33550], image.tag_v2[33922], image.tag_v2[34735]
        entries = {keys[i]: keys[i+3] for i in range(4, len(keys), 4)}
        assert entries[2048] == 4326 and entries[1025] == 2
        self.xcentre = tie[3]-tie[0]*scale[0]
        self.ycentre = tie[4]+tie[1]*scale[1]
        self.sx, self.sy = scale[:2]
        self.mx = self.origin['longitude_scale_m_per_radian']
        self.my = self.origin['latitude_scale_m_per_radian']
        self.lon0, self.lat0 = raw['O01']['lon'], raw['O01']['lat']
        image.close()
        wb = openpyxl.load_workbook(base/'通信链路参数.xlsx', read_only=True, data_only=True)
        rows = list(wb['数据'].values)
        wb.close()
        self.frequency_mhz, self.system_loss_db, self.obstacle_loss_db = (float(rows[i][4]) for i in (2,3,4))
        self.receiver_threshold_dbm = float(rows[5][4])+float(rows[6][4])
        transport = (float(rows[7][4]), float(rows[8][4]))
        access = (float(rows[9][4]), float(rows[10][4]))
        backhaul = (float(rows[11][4]), float(rows[12][4]))
        gateway = (float(rows[13][4]), float(rows[14][4]))
        self.thresholds = {kind:min(a[0],b[0])+a[1]+b[1]-self.system_loss_db-self.receiver_threshold_dbm
                           for kind,a,b in [('direct',transport,gateway),('access',transport,access),('backhaul',backhaul,gateway)]}
        self.gateway = (0.,0.,raw['O01']['elevation_m']+float(rows[15][4]))
        self.fspl_constant = 32.45+20*math.log10(self.frequency_mhz)-60
        self.radii_m = {k:{'los':10**((v-self.fspl_constant)/20),
                             'blocked':10**((v-self.obstacle_loss_db-self.fspl_constant)/20)} for k,v in self.thresholds.items()}
        wb = openpyxl.load_workbook(base/'中继无人机数据.xlsx', read_only=True, data_only=True)
        r = list(wb['数据'].values)[2]
        wb.close()
        self.relay = dict(mass_kg=float(r[4]), cruise_mps=float(r[5]), cruise_kw=float(r[6]), energy_kwh=float(r[7]),
                          reserve_percent=float(r[8]), prepare_s=float(r[9]), setup_s=float(r[10]), turnaround_s=float(r[11]),
                          climb_mps=float(r[12]), descend_mps=float(r[13]), climb_efficiency=float(r[14]),
                          hover_kw=float(r[16]), communication_kw=float(r[17]), max_agl_m=float(r[18]))

    def lonlat(self, x, y):
        return self.lon0+math.degrees(x/self.mx), self.lat0+math.degrees(y/self.my)

    def pixel(self, x, y):
        lon, lat = self.lonlat(x,y)
        return (lon-self.xcentre)/self.sx+.5, (self.ycentre-lat)/self.sy+.5

    def xyz(self, sid, work=True):
        n=self.nodes[sid]
        return n['x'],n['y'],n['z']+(30 if work and sid!='O01' else 0)

    def terrain(self, x, y):
        c,r=self.pixel(x,y)
        r,c=math.floor(r),math.floor(c)
        if not (0<=r<self.dem.shape[0] and 0<=c<self.dem.shape[1]):
            raise ValueError('Coordinate outside DEM')
        value=float(self.dem[r,c])
        if not math.isfinite(value) or value==-32767:
            raise ValueError('Coordinate on DEM nodata')
        return value

    def _cells(self,a,b):
        return set(cells_on_segment(*self.pixel(*a[:2]),*self.pixel(*b[:2])))

    @lru_cache(maxsize=20000)
    def leg(self,i,j):
        a=self.xyz(i) if isinstance(i,str) else tuple(i)
        b=self.xyz(j) if isinstance(j,str) else tuple(j)
        cells=self._cells(a,b)
        if not all(0<=r<self.dem.shape[0] and 0<=c<self.dem.shape[1] for r,c in cells):
            raise ValueError('Flight path outside DEM')
        peak=max(float(self.dem[r,c]) for r,c in cells)
        # For relay endpoints above terrain+50, cruise must also reach the
        # endpoint altitude. This explicit physical-consistency convention
        # supplements the appendix, which otherwise describes node flights.
        altitude=max(peak+50,a[2],b[2])
        return dict(distance_m=math.dist(a[:2],b[:2]),max_terrain_m=peak,cruise_altitude_m=altitude,
                    up_m=altitude-a[2],down_m=altitude-b[2],crossed_cell_count=len(cells))

    @staticmethod
    def _interval_in_cell(p,q,c,r):
        lo,hi=0.,1.
        for start,end,left,right in ((p[0],q[0],c,c+1),(p[1],q[1],r,r+1)):
            d=end-start
            if abs(d)<1e-14:
                if start<left-1e-10 or start>right+1e-10:return None
            else:
                aa,bb=sorted(((left-start)/d,(right-start)/d))
                lo,hi=max(lo,aa),min(hi,bb)
                if lo>hi+1e-10:return None
        return lo,hi

    def blocked(self,a,b):
        p,q=self.pixel(*a[:2]),self.pixel(*b[:2])
        for r,c in set(cells_on_segment(*p,*q)):
            if not (0<=r<self.dem.shape[0] and 0<=c<self.dem.shape[1]):return True
            interval=self._interval_in_cell(p,q,c,r)
            if interval is None:continue
            zmin=min(a[2]+t*(b[2]-a[2]) for t in interval)
            if self.dem[r,c]>=zmin-1e-9:return True
        return False

    @lru_cache(maxsize=300000)
    def _link_cached(self,a,b,kind):
        distance=math.dist(a,b)
        blocked=self.blocked(a,b)
        loss=self.fspl_constant+20*math.log10(max(distance,1e-6))+self.obstacle_loss_db*blocked
        margin=self.thresholds[kind]-loss
        return dict(available=margin>=-1e-9,path_loss_db=loss,margin_db=margin,blocked=blocked,distance_m=distance)

    def link(self,a_xyz,b_xyz,kind='direct'):
        return self._link_cached(tuple(a_xyz),tuple(b_xyz),kind).copy()

    def relay_travel(self,point):
        point=tuple(point)
        ground=self.terrain(*point[:2])
        assert ground<=point[2]<=ground+self.relay['max_agl_m']+1e-7
        leg=self.leg(self.xyz('O01'),point)
        r=self.relay
        outward=leg['up_m']/r['climb_mps']+leg['distance_m']/r['cruise_mps']+leg['down_m']/r['descend_mps']
        inward=leg['down_m']/r['climb_mps']+leg['distance_m']/r['cruise_mps']+leg['up_m']/r['descend_mps']
        energy=2*r['cruise_kw']*leg['distance_m']/r['cruise_mps']/3600 + r['mass_kg']*9.81*(leg['up_m']+leg['down_m'])/(r['climb_efficiency']*3.6e6)
        max_service=(r['energy_kwh']*(1-r['reserve_percent']/100)-energy)*3600/(r['hover_kw']+r['communication_kw'])-r['setup_s']
        return dict(ground_m=ground,agl_m=point[2]-ground,outbound_s=outward,return_s=inward,
                    earliest_service_s=r['prepare_s']+outward+r['setup_s'],travel_energy_kwh=energy,
                    max_service_s=max_service,cruise_altitude_m=leg['cruise_altitude_m'],distance_m=leg['distance_m'])

    @staticmethod
    def _clip(poly,axis,value,keep_greater):
        out=[]
        if not poly:return out
        previous=poly[-1]
        prev_in=(previous[axis]>=value-1e-10) if keep_greater else (previous[axis]<=value+1e-10)
        for current in poly:
            cur_in=(current[axis]>=value-1e-10) if keep_greater else (current[axis]<=value+1e-10)
            if cur_in!=prev_in:
                t=(value-previous[axis])/(current[axis]-previous[axis])
                out.append(tuple(previous[k]+t*(current[k]-previous[k]) for k in range(3)))
            if cur_in:out.append(current)
            previous,prev_in=current,cur_in
        return out

    def fan_clear(self,a,b,anchor):
        """Prove every anchor--p LOS clear for p in closed segment a--b.

        Intersect each DEM pixel prism with the swept triangular LOS surface.
        A positive certificate proves the whole continuum, not time samples.
        """
        vertices=[(*self.pixel(*p[:2]),p[2]) for p in (a,b,anchor)]
        va,vb,vc=vertices
        det=(vb[0]-va[0])*(vc[1]-va[1])-(vb[1]-va[1])*(vc[0]-va[0])
        if abs(det)<1e-9:
            # Projection is a line. The fan is the union of endpoint rays
            # or a triangle on that vertical plane; the lower boundary is
            # among its three edges, hence checking all three is sufficient.
            return not any(self.blocked(x,y) for x,y in ((a,anchor),(b,anchor),(a,b)))
        # Include the neighbouring pixel for an exactly touched lower grid
        # boundary, matching the flight-path supercover convention.
        cmin,cmax=math.floor(min(v[0] for v in vertices)-1e-10),math.floor(max(v[0] for v in vertices))
        rmin,rmax=math.floor(min(v[1] for v in vertices)-1e-10),math.floor(max(v[1] for v in vertices))
        if cmin<0 or rmin<0 or cmax>=self.dem.shape[1] or rmax>=self.dem.shape[0]:return False
        zmin=min(v[2] for v in vertices)
        subset=self.dem[rmin:rmax+1,cmin:cmax+1]
        rows,cols=np.where(subset>=zmin-1e-9)
        for rr,cc in zip(rows,cols):
            r,c=int(rr+rmin),int(cc+cmin)
            poly=vertices
            for axis,value,greater in ((0,c,True),(0,c+1,False),(1,r,True),(1,r+1,False)):
                poly=self._clip(poly,axis,value,greater)
                if not poly:break
            if poly and float(self.dem[r,c])>=min(p[2] for p in poly)-1e-9:return False
        return True

    def segment_certificate(self,a,b,relaypoints=(),min_length_m=.25,max_depth=20):
        """Return continuous coverage certificate; unproven pieces fail closed."""
        a,b=tuple(a),tuple(b)
        anchors=[('G01',self.gateway,'direct')]
        for idx,p in enumerate(relaypoints):
            p=tuple(p)
            if self.link(p,self.gateway,'backhaul')['available']:
                anchors.append((f'R{idx+1}',p,'access'))
        records=[]
        def visit(left,right,t0,t1,depth):
            middle=tuple((x+y)/2 for x,y in zip(left,right))
            for sample,t in ((left,t0),(middle,(t0+t1)/2),(right,t1)):
                if not any(self.link(sample,p,kind)['available'] for _,p,kind in anchors):
                    records.append(dict(t0=t0,t1=t1,source=None,proof='uncovered_witness',witness_t=t,witness=list(sample)))
                    return False
            for name,p,kind in anchors:
                maximum=max(math.dist(left,p),math.dist(right,p))
                if maximum<=self.radii_m[kind]['blocked']-1e-7:
                    records.append(dict(t0=t0,t1=t1,source=name,proof='distance_even_if_blocked',min_margin_db=self.thresholds[kind]-self.fspl_constant-20*math.log10(max(maximum,1e-6))-self.obstacle_loss_db))
                    return True
                if maximum<=self.radii_m[kind]['los']-1e-7 and all(self.link(q,p,kind)['available'] for q in (left,middle,right)) and self.fan_clear(left,right,p):
                    records.append(dict(t0=t0,t1=t1,source=name,proof='DEM_triangle_LOS',min_margin_db=self.thresholds[kind]-self.fspl_constant-20*math.log10(max(maximum,1e-6))))
                    return True
            if depth>=max_depth or math.dist(left,right)<=min_length_m:
                records.append(dict(t0=t0,t1=t1,source=None,proof='unproven',midpoint=list(middle)))
                return False
            tm=(t0+t1)/2
            ok1=visit(left,middle,t0,tm,depth+1)
            ok2=visit(middle,right,tm,t1,depth+1) if ok1 else False
            return ok1 and ok2
        proved=visit(a,b,0.,1.,0)
        return dict(proved=proved,segments=records,unproven_count=sum(r['source'] is None for r in records))

    def route_certificate(self,i,j,relaypoints=()):
        a,b=self.xyz(i),self.xyz(j)
        h=self.leg(i,j)['cruise_altitude_m']
        aa,bb=(a[0],a[1],h),(b[0],b[1],h)
        pieces=[self.segment_certificate(x,y,relaypoints) for x,y in ((a,aa),(aa,bb),(bb,b))]
        return dict(proved=all(p['proved'] for p in pieces),phases=dict(zip(('climb','cruise','descend'),pieces)))

    def candidate_search(self,grid_m=1000):
        services=[k for k in self.nodes if k!='O01']
        direct={s:self.link(self.xyz(s),self.gateway,'direct') for s in services}
        needed=[s for s in services if not direct[s]['available']]
        xy={(n['x'],n['y']) for n in self.nodes.values()}
        for x in np.arange(-6000,6001,grid_m):
            for y in np.arange(-1000,8501,grid_m):xy.add((float(x),float(y)))
        candidates=[]
        for x,y in sorted(xy):
            try:ground=self.terrain(x,y)
            except ValueError:continue
            for agl in (50.,100.,200.,300.):
                p=(x,y,ground+agl)
                bh=self.link(p,self.gateway,'backhaul')
                if not bh['available']:continue
                travel=self.relay_travel(p)
                if travel['max_service_s']<=0:continue
                cover=[s for s in needed if self.link(self.xyz(s),p,'access')['available']]
                if not cover:continue
                margins={s:self.link(self.xyz(s),p,'access')['margin_db'] for s in cover}
                candidates.append(dict(x=x,y=y,z=p[2],lon=self.lonlat(x,y)[0],lat=self.lonlat(x,y)[1],backhaul=bh,
                                       covered_needed_services=cover,access_margins_db=margins,**travel))
        candidates.sort(key=lambda c:(-len(c['covered_needed_services']),c['travel_energy_kwh'],-min(c['access_margins_db'].values())))
        for idx,c in enumerate(candidates):c['id']=f'H{idx+1:04}'
        singles=[c['id'] for c in candidates if len(c['covered_needed_services'])==len(needed)]
        pairs=[]
        all_needed=set(needed)
        # One cheapest representative per identical coverage set suffices
        # for static set-cover feasibility and summed travel-energy ranking.
        unique={}
        for c in sorted(candidates,key=lambda c:c['travel_energy_kwh']):
            unique.setdefault(tuple(c['covered_needed_services']),c)
        representatives=list(unique.values())
        for i,a in enumerate(representatives):
            for b in representatives[i+1:]:
                if set(a['covered_needed_services'])|set(b['covered_needed_services'])>=all_needed:
                    pairs.append((a['travel_energy_kwh']+b['travel_energy_kwh'],a['id'],b['id']))
        pairs.sort()
        return dict(source='original attachments only',thresholds_db=self.thresholds,radii_m=self.radii_m,origin=self.origin,
                    fspl_constant_db=32.45,terrain_model='Piecewise-constant DEM pixels; both sides of touched grid boundaries included.',
                    relay_cruise_convention='max(DEM path maximum + 50 m, both endpoint altitudes); explicit physical-consistency supplement for high relay endpoints.',
                    direct_service_links=direct,needed_services=needed,candidate_count=len(candidates),single_site_ids=singles,
                    unique_coverage_set_count=len(representatives),pair_search='All distinct coverage sets; finite candidate set only.',
                    recommended_pairs=[list(x[1:]) for x in pairs[:20]],candidates=candidates)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--grid-m',type=float,default=1000)
    parser.add_argument('--certify-sites',nargs='*')
    parser.add_argument('--verify-certificates',action='store_true')
    args=parser.parse_args()
    geom=Geometry()
    if args.verify_certificates:
        source=ROOT/'results'/'question3_independent'/'geometry_star_certificates.json'
        data=json.loads(source.read_text(encoding='utf-8'))
        checks=[]
        for site in data['sites']:
            relay=(site['x'],site['y'],site['z'])
            for sid,cert in data['certificates'][site['id']].items():
                if not cert['proved']:continue
                a,b=geom.xyz('O01'),geom.xyz(sid)
                h=geom.leg('O01',sid)['cruise_altitude_m']
                aa,bb=(a[0],a[1],h),(b[0],b[1],h)
                for phase,(left,right) in zip(('climb','cruise','descend'),((a,aa),(aa,bb),(bb,b))):
                    for record in cert['phases'][phase]['segments']:
                        anchor=geom.gateway if record['source']=='G01' else relay
                        kind='direct' if record['source']=='G01' else 'access'
                        for ratio in (0.,.2113248654,.5,.7886751346,1.):
                            t=record['t0']+(record['t1']-record['t0'])*ratio
                            p=tuple(x+t*(y-x) for x,y in zip(left,right))
                            actual=geom.link(p,anchor,kind)
                            checks.append(actual['margin_db'])
                            assert actual['available'],(site['id'],sid,phase,record,t,actual)
        result=dict(check='Point-link cross-check within already proved continuous intervals; not used as the continuous proof.',
                    sampled_links=len(checks),minimum_sampled_margin_db=min(checks),passed=True,
                    star_site_count=len(data['sites']),fspl_constant_db=32.45,
                    thresholds_db=geom.thresholds)
        source.with_name('geometry_verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
        print(json.dumps(result))
        return
    if args.certify_sites is not None:
        source=ROOT/'results'/'question3_independent'/'geometry_candidates.json'
        data=json.loads(source.read_text(encoding='utf-8'))
        chosen=[c for c in data['candidates'] if c['id'] in args.certify_sites]
        result=dict(source='original attachments only',sites=[],direct_roundtrips=[],certificates={})
        for sid in geom.nodes:
            if sid=='O01':continue
            cert=geom.route_certificate('O01',sid)
            if cert['proved']:result['direct_roundtrips'].append(sid)
        output=source.with_name('geometry_star_certificates.json')
        for c in chosen:
            point=(c['x'],c['y'],c['z'])
            certs={sid:geom.route_certificate('O01',sid,[point]) for sid in geom.nodes if sid!='O01'}
            c=dict(c,proved_roundtrips=[s for s,p in certs.items() if p['proved']])
            result['sites'].append(c)
            result['certificates'][c['id']]=certs
            output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            print(c['id'],c['proved_roundtrips'],flush=True)
        return
    result=geom.candidate_search(args.grid_m)
    output=ROOT/'results'/'question3_independent'/'geometry_candidates.json'
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:result[k] for k in ('thresholds_db','radii_m','needed_services','candidate_count','single_site_ids','recommended_pairs')},ensure_ascii=True))
    print(output)


if __name__=='__main__':main()
