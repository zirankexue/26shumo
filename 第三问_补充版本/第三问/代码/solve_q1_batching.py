"""Question 1: exact single-service-area batching under stated energy assumptions.

Run from any directory: python code/solve_q1_batching.py
Dependencies: numpy, Pillow, openpyxl (see requirements_q1.txt).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import itertools
import json
import math
import platform
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if (ROOT / '.runtime_py').is_dir():
    sys.path.insert(0, str(ROOT / '.runtime_py'))

import numpy as np
import openpyxl
import PIL
from PIL import Image

TYPE_ORDER = ('MED', 'WAT', 'FOD', 'HYG')
EPS_ENERGY = 1e-9  # feasibility tolerance, aligned with the peer configuration
EPS_TIME = 1e-6
PAYLOAD_TOLERANCE_KG = 1e-6
ENERGY_COST_SCALE = 10**12
TIME_COST_SCALE = 10**6
GRAVITY_M_S2 = 9.80665
COORDINATE_MODEL = 'WGS84_local_ellipsoidal_plane_O01'


def read_inputs(root):
    base = root / '数据' / '无人机应急物资运输基础数据'
    files = {key: base / name for key, name in {
        'nodes': '调度中心与服务区.xlsx',
        'drones': '运输无人机数据.xlsx',
        'boxes': '物资需求与配送时限.xlsx',
    }.items()}
    files['dem'] = next((root / '数据').rglob('*.tif'))
    files['problem'] = root / '山区洪涝灾害下无人机运输与通信协同优化.docx'
    nodes = {}
    wb = openpyxl.load_workbook(files['nodes'], read_only=True, data_only=True)
    for row in wb['数据'].values:
        if row[0] == 'O01' or (isinstance(row[0], str) and row[0].startswith('S0')):
            nodes[row[0]] = dict(id=row[0], name=row[1], lon=float(row[2]),
                                lat=float(row[3]), elevation_m=float(row[4]))
    wb.close()
    names = ('id', 'name', 'empty_mass_kg', 'payload_kg', 'volume_m3', 'cruise_mps',
             'empty_range_m', 'full_range_m', 'energy_kwh', 'reserve_percent',
             'prepare_s', 'load_per_box_s', 'handover_base_s', 'handover_per_box_s',
             'climb_mps', 'descend_mps', 'climb_efficiency', 'descent_extra_efficiency')
    wb = openpyxl.load_workbook(files['drones'], read_only=True, data_only=True)
    drones = {}
    for row in wb['数据'].iter_rows(min_row=3, max_row=5, values_only=True):
        drone = dict(zip(names, row))
        drone['volume_litre'] = round(drone['volume_m3'] * 1000)
        assert abs(drone['volume_litre'] / 1000 - drone['volume_m3']) < 1e-12
        assert drone['descent_extra_efficiency'] == 0
        drones[drone['id']] = drone
    wb.close()
    wb = openpyxl.load_workbook(files['boxes'], read_only=True, data_only=True)
    boxes = []
    for row in wb['逐箱货箱清单'].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        box = dict(id=row[0], service_id=row[1], type=row[0].split('-')[1],
                   type_name=row[2], mass_kg=float(row[3]), volume_m3=float(row[4]),
                   first_batch=row[5], first_deadline_s=row[6], due_s=row[7], priority=row[8])
        box['volume_litre'] = round(box['volume_m3'] * 1000)
        assert abs(box['volume_litre'] / 1000 - box['volume_m3']) < 1e-12
        boxes.append(box)
    wb.close()
    assert len(boxes) == len({b['id'] for b in boxes}) == 80
    assert set(drones) == {'A', 'B', 'C'} and len(nodes) == 16
    assert math.isclose(sum(b['mass_kg'] for b in boxes), 758)
    assert sum(b['volume_litre'] for b in boxes) == 2011
    attrs = {}
    for typ in TYPE_ORDER:
        values = {(b['mass_kg'], b['volume_litre']) for b in boxes if b['type'] == typ}
        assert len(values) == 1, 'Quantity aggregation requires identical box attributes.'
        attrs[typ] = next(iter(values))
    return files, nodes, drones, boxes, attrs


def cells_on_segment(x0, y0, x1, y1):
    """Supercover of a segment in corner-based pixel coordinates; both sides on ties."""
    cuts = [0.0, 1.0]
    for start, end in ((x0, x1), (y0, y1)):
        if abs(end - start) > 1e-15:
            for boundary in range(math.floor(min(start, end)) + 1, math.ceil(max(start, end))):
                t = (boundary - start) / (end - start)
                if 0 < t < 1:
                    cuts.append(t)
    cuts = sorted(set(cuts))
    points = cuts + [(a + b) / 2 for a, b in zip(cuts, cuts[1:])]
    for t in points:
        x, y = x0 + t * (x1 - x0), y0 + t * (y1 - y0)
        cols, rows = {math.floor(x)}, {math.floor(y)}
        if abs(x - round(x)) < 1e-10:
            cols.update((round(x) - 1, round(x)))
        if abs(y - round(y)) < 1e-10:
            rows.update((round(y) - 1, round(y)))
        yield from itertools.product(rows, cols)


def local_plane(nodes):
    """Fixed WGS84 curvature scales at O01; angles converted to radians."""
    depot = nodes['O01']
    phi = math.radians(depot['lat'])
    a, e2 = 6378137.0, 6.6943799901413165e-3
    w = math.sqrt(1 - e2 * math.sin(phi)**2)
    prime_vertical_radius = a / w
    meridional_radius = a * (1-e2) / w**3
    scale_x, scale_y = prime_vertical_radius * math.cos(phi), meridional_radius
    coordinates = {sid: dict(x_m=scale_x * math.radians(node['lon']-depot['lon']),
                            y_m=scale_y * math.radians(node['lat']-depot['lat']))
                   for sid, node in nodes.items()}
    metadata = dict(model=COORDINATE_MODEL, origin_id='O01',
                    origin_lon_degree=depot['lon'], origin_lat_degree=depot['lat'],
                    ellipsoid='WGS84', semi_major_axis_m=a, eccentricity_squared=e2,
                    prime_vertical_radius_m=prime_vertical_radius,
                    meridional_radius_m=meridional_radius,
                    longitude_scale_m_per_radian=scale_x, latitude_scale_m_per_radian=scale_y,
                    formula='x=N0*cos(phi0)*radians(lon-lon0); y=M0*radians(lat-lat0)',
                    node_coordinates=coordinates)
    return coordinates, metadata


def compute_routes(nodes, dem_path):
    image = Image.open(dem_path)
    dem = np.asarray(image)
    scale, tie, keys = image.tag_v2[33550], image.tag_v2[33922], image.tag_v2[34735]
    entries = {keys[i]: keys[i+3] for i in range(4, len(keys), 4)}
    assert entries[2048] == 4326 and entries[1025] == 2, 'Expected WGS84 PixelIsPoint.'
    xcentre = tie[3] - tie[0] * scale[0]
    ycentre = tie[4] + tie[1] * scale[1]
    coordinates, plane_meta = local_plane(nodes)
    depot = nodes['O01']
    x0 = (depot['lon']-xcentre)/scale[0]+0.5
    y0 = (ycentre-depot['lat'])/scale[1]+0.5
    routes = {}
    for sid, node in sorted(nodes.items()):
        if sid == 'O01':
            continue
        xy = coordinates[sid]
        distance = math.hypot(xy['x_m'], xy['y_m'])
        # This fixed linear coordinate map sends a local straight route to an
        # exact lon/lat straight segment. Enumerate every touched DEM cell.
        x1 = (node['lon']-xcentre)/scale[0]+0.5
        y1 = (ycentre-node['lat'])/scale[1]+0.5
        cells = set(cells_on_segment(x0, y0, x1, y1))
        ordered = sorted(cells)
        assert all(0 <= r < dem.shape[0] and 0 <= c < dem.shape[1] for r, c in ordered)
        heights = np.array([dem[r,c] for r,c in ordered], dtype=float)
        assert np.isfinite(heights).all() and not (heights == -32767).any()
        peak = float(heights.max())
        altitude = peak + 50
        up_out, up_back = altitude-depot['elevation_m'], altitude-node['elevation_m']-30
        assert up_out >= 0 and up_back >= 0, 'Cruise altitude conflicts with node operation height.'
        routes[sid] = dict(service_id=sid, distance_m=distance, max_terrain_m=peak,
                           cruise_altitude_m=altitude, outbound_climb_m=up_out,
                           return_climb_m=up_back, crossed_cell_count=len(ordered),
                           peak_cells=[list(ordered[j]) for j in np.flatnonzero(heights == peak)],
                           crossed_cells=[list(pair) for pair in ordered])
    metadata = dict(crs='EPSG:4326', raster_type='PixelIsPoint', shape=list(dem.shape),
                    centre_origin=[xcentre, ycentre], pixel_size_degree=list(scale[:2]),
                    nodata=-32767, distance_crs=COORDINATE_MODEL, local_plane=plane_meta,
                    route_geometry='Straight segment in fixed local plane and lon/lat; exact grid crossings.',
                    grid_boundary_rule='Include cells on both sides of a touched boundary.')
    return routes, metadata


def energy(drone, route, mass):
    equivalent_range = drone['empty_range_m'] - (drone['empty_range_m'] - drone['full_range_m']) * (mass/drone['payload_kg'])**1.5
    horizontal_out = drone['energy_kwh'] * route['distance_m'] / equivalent_range
    horizontal_back = drone['energy_kwh'] * route['distance_m'] / drone['empty_range_m']
    factor = GRAVITY_M_S2 / (drone['climb_efficiency'] * 3.6e6)
    climb_out = (drone['empty_mass_kg']+mass) * route['outbound_climb_m'] * factor
    climb_back = drone['empty_mass_kg'] * route['return_climb_m'] * factor
    return dict(horizontal_out_kwh=horizontal_out, horizontal_return_kwh=horizontal_back,
                climb_out_kwh=climb_out, climb_return_kwh=climb_back,
                energy_kwh=horizontal_out+horizontal_back+climb_out+climb_back)


def times(drone, route, count):
    flight = 2*route['distance_m']/drone['cruise_mps'] + (route['outbound_climb_m']+route['return_climb_m'])*(1/drone['climb_mps']+1/drone['descend_mps'])
    return dict(flight_time_s=flight, work_time_s=flight+drone['prepare_s']+
                count*drone['load_per_box_s']+drone['handover_base_s']+count*drone['handover_per_box_s'])


def safe_payload(drone, route):
    budget = drone['energy_kwh']*(1-drone['reserve_percent']/100)
    if energy(drone, route, 0)['energy_kwh'] > budget:
        return dict(safe_payload_kg=None, limiting_factor='empty_roundtrip_infeasible')
    if energy(drone, route, drone['payload_kg'])['energy_kwh'] <= budget:
        return dict(safe_payload_kg=drone['payload_kg'], limiting_factor='rated_payload')
    low, high = 0.0, float(drone['payload_kg'])
    while high-low > PAYLOAD_TOLERANCE_KG:
        mid = (low+high)/2
        if energy(drone, route, mid)['energy_kwh'] <= budget:
            low = mid
        else:
            high = mid
    return dict(safe_payload_kg=low, limiting_factor='energy_reserve')


def better(value, incumbent):
    """Compare sums of per-batch integer costs, never rounded running totals."""
    return incumbent is None or value < incumbent


def batch_cost(batch):
    return (1, round(batch['energy_kwh']*ENERGY_COST_SCALE),
            round(batch['work_time_s']*TIME_COST_SCALE))


def solve_service(sid, boxes, drones, route, attrs):
    demand = tuple(sum(b['type'] == typ for b in boxes) for typ in TYPE_ORDER)
    candidates, capacities = [], {}
    for gid, drone in drones.items():
        capacities[gid] = safe_payload(drone, route)
        for combo in itertools.product(*(range(n+1) for n in demand)):
            count = sum(combo)
            if not count:
                continue
            mass = sum(attrs[t][0]*c for t,c in zip(TYPE_ORDER,combo))
            litres = sum(attrs[t][1]*c for t,c in zip(TYPE_ORDER,combo))
            if mass > drone['payload_kg'] or litres > drone['volume_litre']:
                continue
            components = energy(drone, route, mass)
            budget = drone['energy_kwh']*(1-drone['reserve_percent']/100)
            if components['energy_kwh'] > budget+EPS_ENERGY:
                continue
            candidates.append(dict(service_id=sid, model=gid, counts=list(combo), box_count=count,
                                   mass_kg=mass, volume_m3=litres/1000, **components,
                                   **times(drone, route, count),
                                   return_soc_percent=100*(1-components['energy_kwh']/drone['energy_kwh']),
                                   payload_limit_kg=drone['payload_kg'], volume_limit_m3=drone['volume_m3'],
                                   energy_budget_kwh=budget, reserve_percent=drone['reserve_percent']))
    candidates.sort(key=lambda p:(p['counts'],p['model']))
    states = sorted(itertools.product(*(range(n+1) for n in demand)), key=lambda s:(sum(s),s))
    dp, parent = {(0,0,0,0):(0,0,0)}, {}
    for state in states[1:]:
        best, pred = None, None
        for idx, p in enumerate(candidates):
            remaining = tuple(s-c for s,c in zip(state,p['counts']))
            if min(remaining) < 0 or remaining not in dp:
                continue
            previous = dp[remaining]
            value = tuple(a+b for a,b in zip(previous,batch_cost(p)))
            if better(value,best):
                best, pred = value, (remaining,idx)
        if best is not None:
            dp[state], parent[state] = best,pred
    if demand not in dp:
        raise ValueError(f'{sid}: no feasible batching covers every box.')
    selected, state = [], demand
    while any(state):
        previous, idx = parent[state]
        selected.append(dict(candidates[idx]))
        state = previous
    selected.sort(key=lambda p:(p['model'],p['counts']))
    available = {t:sorted(b['id'] for b in boxes if b['type']==t) for t in TYPE_ORDER}
    for p in selected:
        p['box_ids'] = []
        for typ,count in zip(TYPE_ORDER,p['counts']):
            p['box_ids'].extend(available[typ][:count])
            available[typ] = available[typ][count:]
    assert not any(available.values())
    mass_total = sum(b['mass_kg'] for b in boxes)
    volume_total = sum(b['volume_litre'] for b in boxes)
    simple_lower_bound = max(1, math.ceil(mass_total/max(d['payload_kg'] for d in drones.values())),
                             math.ceil(volume_total/max(d['volume_litre'] for d in drones.values())))
    summary = dict(service_id=sid, demand=list(demand), box_count=len(boxes), mass_kg=mass_total,
                   volume_m3=volume_total/1000, flight_count=dp[demand][0], energy_kwh=math.fsum(p['energy_kwh'] for p in selected),
                   work_time_s=math.fsum(p['work_time_s'] for p in selected), flight_time_s=math.fsum(p['flight_time_s'] for p in selected),
                   model_counts=dict(Counter(p['model'] for p in selected)),
                   minimum_soc_percent=min(p['return_soc_percent'] for p in selected),
                   state_count=len(states), feasible_candidate_count=len(candidates),
                   simple_flight_lower_bound=simple_lower_bound,
                   safe_payloads=capacities)
    return selected, summary, candidates


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT/'results'/'question1_batching')
    args = parser.parse_args()
    files, nodes, drones, boxes, attrs = read_inputs(ROOT)
    routes, dem_meta = compute_routes(nodes, files['dem'])
    batches, summaries, candidate_map = [], [], {}
    for sid, route in routes.items():
        selected, summary, candidates = solve_service(sid, [b for b in boxes if b['service_id']==sid], drones, route, attrs)
        batches.extend(selected)
        summaries.append(summary)
        candidate_map[sid] = candidates
    for n,p in enumerate(batches,1):
        p['batch_id'] = f'Q1-{n:03d}'
    delivered = [bid for p in batches for bid in p['box_ids']]
    assert Counter(delivered) == Counter(b['id'] for b in boxes)
    assert all(p['mass_kg']<=p['payload_limit_kg'] and p['volume_m3']<=p['volume_limit_m3']+1e-12
               and p['energy_kwh']<=p['energy_budget_kwh']+EPS_ENERGY for p in batches)
    totals = dict(flight_count=len(batches), box_count=len(delivered), mass_kg=sum(p['mass_kg'] for p in batches),
                  volume_m3=math.fsum(p['volume_m3'] for p in batches), energy_kwh=math.fsum(p['energy_kwh'] for p in batches),
                  flight_time_s=math.fsum(p['flight_time_s'] for p in batches), work_time_s=math.fsum(p['work_time_s'] for p in batches),
                  model_counts=dict(Counter(p['model'] for p in batches)),
                  minimum_soc_percent=min(p['return_soc_percent'] for p in batches),
                  simple_flight_lower_bound=sum(s['simple_flight_lower_bound'] for s in summaries))
    solution = dict(created_utc=datetime.now(timezone.utc).isoformat(),
        objective=['minimum_flight_count','minimum_energy_kwh','minimum_cumulative_work_time_s'],
        scope='Q1 single-service-area batching, no physical drone/battery scheduling or delivery deadlines',
        assumptions=['Horizontal energy = available battery energy * horizontal distance / stated equivalent range.',
                     f'Extra ascent energy = total mass * {GRAVITY_M_S2} * ascent / (efficiency * 3.6e6).',
                     'All goods delivered at the only service area; return payload is zero.',
                     'Preparation, loading, flight and handover durations are summed sequentially.',
                     'Fixed WGS84 curvature scales at O01 define local planar distances; all cells touched by the straight lon/lat route are included without DEM resampling.'],
        physical_parameters=dict(gravity_m_s2=GRAVITY_M_S2),
        solver_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        numeric_tolerances=dict(energy_kwh=EPS_ENERGY,time_s=EPS_TIME,payload_bisection_kg=PAYLOAD_TOLERANCE_KG,
                                objective_energy_scale=ENERGY_COST_SCALE,objective_time_scale=TIME_COST_SCALE,
                                comparison='Sum independently rounded per-batch integer costs before lexicographic comparison; report unrounded costs.'),
        versions=dict(python=platform.python_version(),numpy=np.__version__,Pillow=PIL.__version__,
                      openpyxl=openpyxl.__version__),
        sources={k:dict(path=str(p.relative_to(ROOT)),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for k,p in files.items()},
        type_order=list(TYPE_ORDER), nodes=nodes, drones=drones, boxes=boxes,dem_metadata=dem_meta,
        routes=routes, service_summary=summaries, batches=batches, totals=totals,
        checks=dict(all_box_ids_exactly_once=True,all_batch_constraints_pass=True))
    args.output.mkdir(parents=True,exist_ok=True)
    (args.output/'solution.json').write_text(json.dumps(solution,ensure_ascii=False,indent=2),encoding='utf-8')
    (args.output/'candidates.json').write_text(json.dumps(candidate_map,ensure_ascii=False,indent=2),encoding='utf-8')
    headers=['架次编号','服务区编号','机型编号','货箱编号列表','总质量（kg）','总体积（m³）','往返时间（s）','架次能耗（kWh）','返航SOC（%）']
    with (args.output/'Q1_单点组批.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(headers)
        for p in batches:
            writer.writerow([p['batch_id'],p['service_id'],p['model'],';'.join(p['box_ids']),p['mass_kg'],p['volume_m3'],p['flight_time_s'],p['energy_kwh'],p['return_soc_percent']])
    with (args.output/'Q1_局部椭球距离与航线.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['服务区编号', '服务区名称', '单程水平距离_m', '往返水平距离_m',
                         '沿线最高高程_m', '巡航海拔_m', '去程爬升_m', '返程爬升_m'])
        for sid, route in routes.items():
            writer.writerow([sid, nodes[sid]['name'], route['distance_m'], 2*route['distance_m'],
                             route['max_terrain_m'], route['cruise_altitude_m'],
                             route['outbound_climb_m'], route['return_climb_m']])
    with (args.output/'Q1_局部椭球有向航段.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['start', 'end', 'distance', 'cruise_altitude', 'climb', 'descent'])
        for sid, route in routes.items():
            for start, end, climb, descent in [('O01', sid, route['outbound_climb_m'], route['return_climb_m']),
                                                (sid, 'O01', route['return_climb_m'], route['outbound_climb_m'])]:
                writer.writerow([start, end, route['distance_m'], route['cruise_altitude_m'], climb, descent])
    print(json.dumps(totals,ensure_ascii=False,indent=2))
    for p in batches:
        print(p['batch_id'],p['service_id'],p['model'],p['counts'],p['mass_kg'],p['volume_m3'],round(p['energy_kwh'],6),round(p['return_soc_percent'],4))


if __name__ == '__main__':
    main()
