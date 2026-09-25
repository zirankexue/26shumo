"""Independently validate Q3 under the senior Q2 physical/time convention.

Read original attachments, never call the senior RouteFactory or Q3 solver.
Only the separate Geometry DEM/radio primitives are shared. Millisecond ceilings
are numerical waiting, not additional flight energy in the inherited model.

Usage: python code/validate_q3_senior.py --solution PATH --output PATH
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
import os
from pathlib import Path
import re

import openpyxl
from q3_geometry import Geometry

ROOT = Path(__file__).resolve().parents[1]
TOL = 1e-6
GRAVITY = 9.80665
EXPECTED_PHYSICS = dict(gravity_m_s2=GRAVITY, clearance_m=50.0,
                        service_height_m=30.0, reserve=0.2)
WAIT_NAMES = {'量化等待', '取整等待', '毫秒等待', 'stationary', 'rounding_wait'}


def ticks(seconds):
    return math.ceil(seconds * 1000 - 1e-9)


def quantized(seconds):
    return ticks(seconds) / 1000


def recharge_seconds(consumed, capacity, full_seconds):
    missing = consumed / capacity
    return full_seconds * (min(missing, .1) * 3.5 + max(missing - .1, 0) * (.65 / .9))


def read_originals_senior():
    base = ROOT / '数据' / '无人机应急物资运输基础数据'
    book = openpyxl.load_workbook(base / '运输无人机数据.xlsx', read_only=True, data_only=True)
    rows = list(book['数据'].values)
    book.close()
    names = ('id', 'name', 'empty', 'payload', 'volume', 'speed', 'empty_range',
             'full_range', 'capacity', 'reserve', 'prepare', 'load', 'handover',
             'per_box', 'up_speed', 'down_speed', 'efficiency', 'descent_efficiency')
    models = {r[0]: dict(zip(names, r)) for r in rows[2:5]}
    fleet = {r[0]: r[1] for r in rows if isinstance(r[0], str) and r[0].startswith('U')}
    batteries = {r[0]: dict(count=int(r[1]), full_s=float(r[2])) for r in rows[19:22]}
    book = openpyxl.load_workbook(base / '物资需求与配送时限.xlsx', read_only=True, data_only=True)
    boxes = {}
    for r in book['逐箱货箱清单'].iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            boxes[r[0]] = dict(id=r[0], service=r[1], kind=r[0].split('-')[1],
                medical=r[2] == '医疗物资', mass=float(r[3]), volume=float(r[4]),
                first=r[5] == '是', first_deadline=float(r[6]) if r[5] == '是' else None,
                due=float(r[7]), priority=float(r[8]))
    book.close()
    book = openpyxl.load_workbook(base / '中继无人机数据.xlsx', read_only=True, data_only=True)
    rows = list(book['数据'].values)
    book.close()
    r = rows[2]
    relay = dict(mass=r[4], speed=r[5], cruise_kw=r[6], capacity=r[7], reserve=r[8],
        prepare=r[9], setup=r[10], turnaround=r[11], up_speed=r[12], down_speed=r[13],
        efficiency=r[14], hover_kw=r[16], comm_kw=r[17], max_agl=r[18])
    relay_fleet = {r[0] for r in rows if isinstance(r[0], str) and r[0].startswith('R0')}
    return models, fleet, batteries, boxes, relay, relay_fleet, int(rows[11][1]), float(rows[11][2])


class Audit:
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.checks = 0
        self.maximum_errors = defaultdict(float)

    def check(self, condition, message):
        self.checks += 1
        if not condition:
            self.errors.append(message)
        return bool(condition)

    def close(self, actual, expected, key, where, tolerance=TOL):
        error = abs(actual - expected)
        self.maximum_errors[key] = max(self.maximum_errors[key], error)
        return self.check(math.isfinite(error) and error <= tolerance,
                          f'{where}: {key} expected {expected}, got {actual}')

    def grid(self, number, where):
        self.close(number, round(number * 1000) / 1000, 'millisecond_grid_s', where, 1e-7)

    def overlaps(self, entries, name):
        grouped = defaultdict(list)
        for key, start, end, sid in entries:
            self.check(end >= start - TOL, f'{name} {key}: negative interval for {sid}')
            grouped[key].append((start, end, sid))
        for key, intervals in grouped.items():
            intervals.sort()
            for first, second in zip(intervals, intervals[1:]):
                self.check(second[0] >= first[1] - TOL,
                           f'{name} {key} overlap: {first} / {second}')
        return {key: len(value) for key, value in sorted(grouped.items())}

    def report(self, solution, **extra):
        return dict(passed=not self.errors, solution=str(solution), check_count=self.checks,
            error_count=len(self.errors), errors=self.errors, warnings=self.warnings,
            maximum_numeric_errors=dict(self.maximum_errors), **extra)


def schema_check(data, audit):
    """Fail with a field path instead of an unexplained KeyError."""
    def required(obj, fields, where):
        if not audit.check(isinstance(obj, dict), where + ': expected object'):
            return False
        missing = [key for key in fields if key not in obj]
        return audit.check(not missing, f'{where}: missing fields {missing}')

    required(data, ('feasible', 'transport', 'relay', 'deliveries', 'metrics', 'physics', 'senior_source'), 'solution')
    if audit.errors:
        return False
    for name in ('transport', 'relay', 'deliveries'):
        audit.check(isinstance(data[name], list), name + ': expected array')
    audit.check(isinstance(data['metrics'], dict), 'metrics: expected object')
    required(data['physics'], EXPECTED_PHYSICS, 'physics')
    audit.check(isinstance(data['senior_source'], (str, dict)) and bool(data['senior_source']),
                'senior_source: expected nonempty provenance string or object')
    if audit.errors:
        return False
    audit.check(bool(data['transport']), 'transport: empty schedule')
    for index, row in enumerate(data['transport']):
        where = f'transport[{index}]'
        if not required(row, ('id', 'candidate_id', 'drone', 'model', 'battery', 'boxes', 'order',
            'mass_kg', 'volume_m3', 'energy_kwh', 'duration_s', 'flight_start_s', 'start_s', 'return_s',
            'resource_end_s', 'charge_s', 'battery_ready_s', 'delivery', 'delivery_order', 'stages',
            'communication'), where):
            continue
        for name in ('boxes', 'order', 'stages', 'communication'):
            audit.check(isinstance(row[name], list), where + '.' + name + ': expected array')
        audit.check(isinstance(row['delivery'], dict), where + '.delivery: expected object')
        audit.check(isinstance(row['delivery_order'], dict), where + '.delivery_order: expected service-to-box-order object')
        if isinstance(row['stages'], list):
            for j, stage in enumerate(row['stages']):
                required(stage, ('kind', 'leg', 'a', 'b', 'start', 'end'), f'{where}.stages[{j}]')
        if isinstance(row['communication'], list):
            for j, part in enumerate(row['communication']):
                required(part, ('kind', 'leg', 'a', 'b', 'start', 'end', 'relay'), f'{where}.communication[{j}]')
                if isinstance(part, dict) and part.get('relay'):
                    audit.check('station' in part or 'station' in row,
                                f'{where}.communication[{j}]: relay station missing')
    for index, row in enumerate(data['relay']):
        where = f'relay[{index}]'
        if required(row, ('id', 'station', 'station_data', 'drone', 'component', 'start_s',
            'service_start_s', 'service_end_s', 'return_s', 'resource_end_s', 'energy_kwh', 'soc_percent'), where):
            required(row['station_data'], ('x', 'y', 'z'), where + '.station_data')
    for index, row in enumerate(data['deliveries']):
        required(row, ('box', 'sortie', 'service', 'time_s'), f'deliveries[{index}]')
    def finite_values(value, where):
        if isinstance(value, dict):
            for key, child in value.items():
                finite_values(child, where + '.' + str(key))
        elif isinstance(value, list):
            for i, child in enumerate(value):
                finite_values(child, f'{where}[{i}]')
        elif isinstance(value, float):
            audit.check(math.isfinite(value), where + ': non-finite number')
    finite_values(data, 'solution')
    return not audit.errors


def canonical_battery(name, model, count, audit, sid):
    match = re.fullmatch(rf'(?:{re.escape(model)}B|BAT-{re.escape(model)}-)(\d+)', name)
    valid = match is not None and 1 <= int(match.group(1)) <= count
    audit.check(valid, sid + ': battery incompatible or outside inventory: ' + name)
    return f'{model}B{int(match.group(1)):02}' if valid else name


def normalized_kind(kind):
    return '量化等待' if kind in WAIT_NAMES else kind


def recompute_transport_senior(data, geo, originals, audit):
    models, fleet, batteries, boxes = originals[:4]
    physical = {}
    deliveries = {}
    air = []
    battery_intervals = []
    hard_limits = []
    waits = []
    gaps = []
    for sortie in data['transport']:
        sid, model = sortie['id'], sortie['model']
        if not audit.check(model in models, sid + ': unknown model'):
            continue
        if not audit.check(all(b in boxes for b in sortie['boxes']), sid + ': unknown box'):
            continue
        g = models[model]
        assigned = [boxes[b] for b in sortie['boxes']]
        mass = math.fsum(b['mass'] for b in assigned)
        volume = math.fsum(b['volume'] for b in assigned)
        audit.check(bool(assigned), sid + ': empty transport sortie')
        audit.check(mass <= g['payload'] + TOL, sid + ': payload exceeded')
        audit.check(volume <= g['volume'] + TOL, sid + ': volume exceeded')
        audit.close(sortie['mass_kg'], mass, 'mass_kg', sid)
        audit.close(sortie['volume_m3'], volume, 'volume_m3', sid)
        visit_set = {b['service'] for b in assigned}
        audit.check(set(sortie['order']) == visit_set, sid + ': visit set mismatch')
        audit.check(len(sortie['order']) == len(set(sortie['order'])), sid + ': repeated service visit')
        if not audit.check(all(v in geo.nodes and v != 'O01' for v in sortie['order']), sid + ': invalid service node'):
            continue
        audit.check(fleet.get(sortie['drone']) == model, sid + ': aircraft/model mismatch')
        battery = canonical_battery(sortie['battery'], model, batteries[model]['count'], audit, sid)
        audit.check(set(sortie['delivery_order']) == visit_set, sid + ': delivery_order service set mismatch')
        audit.check(set(sortie['delivery']) == set(sortie['boxes']), sid + ': delivery offset set mismatch')
        audit.check(bool(sortie['candidate_id']), sid + ': empty candidate id')
        candidate_key = (model, tuple(sorted(sortie['boxes'])), tuple(sortie['order']))
        candidate_id = hashlib.sha256(repr(candidate_key).encode()).hexdigest()[:16]
        audit.check(sortie['candidate_id'] == candidate_id, sid + ': candidate id does not match cargo/model/visits')
        for field in ('start_s', 'return_s', 'duration_s', 'flight_start_s', 'resource_end_s', 'battery_ready_s'):
            audit.grid(sortie[field], sid + '.' + field)
        audit.check(sortie['start_s'] >= -TOL, sid + ': starts before zero')
        initial = g['prepare'] + len(assigned) * g['load']
        time = quantized(initial)
        flight_start = time
        physical_time = initial
        cargo = mass
        energy = 0.
        last = 'O01'
        stages = []
        legs = []
        sortie_wait = 0.

        def stage(kind, a, b, duration, pair):
            nonlocal time
            stages.append(dict(kind=kind, a=list(a), b=list(b), start=time,
                               end=time + duration, leg=list(pair)))
            time += duration

        for nxt in [*sortie['order'], 'O01']:
            leg = geo.leg(last, nxt)
            a, b = geo.xyz(last), geo.xyz(nxt)
            height = max(leg['max_terrain_m'] + 50., a[2], b[2])
            up, down, distance = height - a[2], height - b[2], leg['distance_m']
            equivalent = g['empty_range'] - (g['empty_range'] - g['full_range']) * (max(0., cargo) / g['payload']) ** 1.5
            horizontal = g['capacity'] * distance / equivalent
            ascent = (g['empty'] + cargo) * GRAVITY * up / (g['efficiency'] * 3.6e6)
            energy += horizontal + ascent
            duration = up / g['up_speed'] + distance / g['speed'] + down / g['down_speed']
            leg_start = time
            leg_end = leg_start + quantized(duration)
            aa, bb = (a[0], a[1], height), (b[0], b[1], height)
            stage('爬升', a, aa, up / g['up_speed'], (last, nxt))
            stage('巡航', aa, bb, distance / g['speed'], (last, nxt))
            stage('下降', bb, b, down / g['down_speed'], (last, nxt))
            wait = leg_end - time
            audit.check(-1e-9 <= wait < .001 + 1e-9, sid + ': invalid millisecond leg padding')
            if wait > 1e-10:
                stage('量化等待', b, b, wait, (last, nxt))
            time = leg_end
            sortie_wait += max(0., wait)
            physical_time += duration
            waits.append(dict(sortie=sid, leg=[last, nxt], start_s=sortie['start_s'] + leg_end - max(0., wait),
                              end_s=sortie['start_s'] + leg_end, wait_s=max(0., wait)))
            legs.append(dict(start=last, end=nxt, remaining_cargo_kg=cargo, distance_m=distance,
                cruise_altitude_m=height, horizontal_energy_kwh=horizontal, ascent_energy_kwh=ascent,
                physical_flight_s=duration, scheduled_flight_s=quantized(duration), quantization_wait_s=max(0., wait)))
            if nxt != 'O01':
                local = [v for v in assigned if v['service'] == nxt]
                order = sortie['delivery_order'].get(nxt, [])
                valid_order = isinstance(order, list) and Counter(order) == Counter(v['id'] for v in local)
                audit.check(valid_order, sid + ': invalid per-box delivery order at ' + nxt)
                base_end = time + quantized(g['handover'])
                service_duration = quantized(g['handover']) + len(local) * quantized(g['per_box'])
                if valid_order:
                    for rank, bid in enumerate(order, 1):
                        offset = base_end + rank * quantized(g['per_box'])
                        absolute = sortie['start_s'] + offset
                        box = boxes[bid]
                        deliveries[bid] = dict(time_s=absolute, sortie=sid, service=nxt, rank=rank)
                        if bid in sortie['delivery']:
                            audit.close(sortie['delivery'][bid], offset, 'delivery_offset_s', sid + '/' + bid)
                        limits = []
                        if box['medical']:
                            limits.append(('medical_due', box['due']))
                        if box['first']:
                            limits.append(('first_deadline', box['first_deadline']))
                        for rule, limit in limits:
                            slack = limit - absolute
                            hard_limits.append(dict(box=bid, rule=rule, deadline_s=limit, delivery_s=absolute, slack_s=slack))
                            audit.check(slack >= -TOL, f'{sid}/{bid}: {rule} violated by {-slack} s')
                stage('投送', b, b, service_duration, (nxt, nxt))
                physical_time += g['handover'] + len(local) * g['per_box']
                cargo -= math.fsum(v['mass'] for v in local)
            last = nxt
        audit.close(cargo, 0., 'final_payload_kg', sid)
        for field, expected in [('energy_kwh', energy), ('duration_s', time), ('flight_start_s', flight_start),
                                ('return_s', sortie['start_s'] + time)]:
            audit.close(sortie[field], expected, field, sid)
        audit.check(sortie['resource_end_s'] >= sortie['return_s'] - TOL, sid + ': aircraft released before return')
        soc = 100 * (1 - energy / g['capacity'])
        audit.check(soc >= 20. - TOL and energy <= .8 * g['capacity'] + 1e-9, sid + ': return SOC below 20%')
        audit.close(g['reserve'], 20., 'source_reserve_percent', sid)
        recharge = recharge_seconds(energy, g['capacity'], batteries[model]['full_s'])
        audit.close(sortie['charge_s'], quantized(recharge), 'charge_s', sid)
        audit.check(sortie['battery_ready_s'] >= sortie['return_s'] + quantized(recharge) - TOL,
                    sid + ': battery reused before full millisecond-ceiled recharge')
        air.append((sortie['drone'], sortie['start_s'], sortie['resource_end_s'], sid))
        battery_intervals.append((battery, sortie['start_s'], sortie['battery_ready_s'], sid))
        supplied_stages = sortie['stages']
        audit.check(len(stages) == len(supplied_stages), sid + ': stage count mismatch')
        previous = flight_start
        previous_point = geo.xyz('O01')
        for index, actual in enumerate(supplied_stages):
            audit.close(actual['start'], previous, 'stage_continuity_s', sid)
            audit.check(actual['end'] >= actual['start'] - TOL, sid + ': negative stage duration')
            audit.close(math.dist(actual['a'], previous_point), 0., 'stage_position_continuity_m', sid)
            previous, previous_point = actual['end'], actual['b']
            if index >= len(stages):
                continue
            expected = stages[index]
            audit.check(normalized_kind(actual['kind']) == expected['kind'] and actual['leg'] == expected['leg'],
                        sid + ': stage identity mismatch at ' + str(index))
            for field in ('start', 'end'):
                audit.close(actual[field], expected[field], 'stage_' + field + '_s', sid)
            for field in ('a', 'b'):
                audit.close(math.dist(actual[field], expected[field]), 0., 'stage_xyz_m', sid)
        audit.close(previous, time, 'final_stage_end_s', sid)
        physical[sid] = dict(stages=stages, legs=legs, energy_kwh=energy, soc_percent=soc,
                            return_s=sortie['start_s'] + time)
        gaps.append(dict(sortie=sid, physical_operation_s=physical_time, quantized_operation_s=time,
                         total_numerical_wait_s=time - physical_time, flight_quantization_wait_s=sortie_wait,
                         physical_charge_s=recharge, charge_quantization_wait_s=quantized(recharge) - recharge,
                         reported_duration_error_s=sortie['duration_s'] - time,
                         reported_energy_error_kwh=sortie['energy_kwh'] - energy))
    return dict(physical=physical, deliveries=deliveries, air=air, batteries=battery_intervals,
                hard_limits=hard_limits, waits=waits, gaps=gaps)


def recompute_relay_senior(data, geo, originals, audit):
    relay, fleet, component_count, full_charge = originals[4:]
    relay_map = {}
    air = []
    components = []
    recomputed = []
    for sortie in data['relay']:
        sid, station = sortie['id'], sortie['station_data']
        point = (station['x'], station['y'], station['z'])
        origin = geo.xyz('O01')
        try:
            ground = geo.terrain(*point[:2])
            leg = geo.leg(origin, point)
        except (ValueError, IndexError) as error:
            audit.check(False, sid + ': relay terrain/flight path invalid: ' + str(error))
            continue
        agl = point[2] - ground
        audit.check(-TOL <= agl <= relay['max_agl'] + TOL, sid + ': relay AGL outside [0, 300] m')
        height = max(leg['max_terrain_m'] + 50., origin[2], point[2])
        up, down = height - origin[2], height - point[2]
        outbound = up / relay['up_speed'] + leg['distance_m'] / relay['speed'] + down / relay['down_speed']
        inward = down / relay['up_speed'] + leg['distance_m'] / relay['speed'] + up / relay['down_speed']
        travel = 2 * leg['distance_m'] / relay['speed'] * relay['cruise_kw'] / 3600
        travel += GRAVITY * relay['mass'] * (up + down) / (relay['efficiency'] * 3.6e6)
        wait = sortie['service_start_s'] - sortie['start_s'] - relay['prepare'] - outbound - relay['setup']
        audit.check(sortie['start_s'] >= -TOL, sid + ': relay starts before zero')
        audit.check(wait >= -TOL, sid + ': relay not at station or setup incomplete when service begins')
        audit.check(sortie['service_end_s'] >= sortie['service_start_s'] - TOL, sid + ': negative service window')
        for field in ('start_s', 'service_start_s', 'service_end_s', 'resource_end_s'):
            audit.grid(sortie[field], sid + '.' + field)
        duration = sortie['service_end_s'] - sortie['service_start_s']
        energy = travel + (duration + relay['setup']) * (relay['hover_kw'] + relay['comm_kw']) / 3600
        soc = 100 * (1 - energy / relay['capacity'])
        physical_return = sortie['service_end_s'] + inward
        padded_return = sortie['service_end_s'] + quantized(inward)
        audit.close(sortie['energy_kwh'], energy, 'relay_energy_kwh', sid)
        audit.close(sortie['soc_percent'], soc, 'relay_soc_percent', sid)
        audit.check(soc >= relay['reserve'] - TOL, sid + ': relay return SOC below reserve')
        audit.close(sortie['return_s'], padded_return, 'relay_return_s', sid)
        audit.grid(sortie['return_s'], sid + '.return_s')
        audit.check(sortie['resource_end_s'] >= padded_return + relay['turnaround'] - TOL,
                    sid + ': relay 300-second turnaround or return quantization omitted')
        if 'ret' in sortie:
            audit.close(sortie['ret'], padded_return, 'relay_ret_s', sid)
        for field, expected in [('outbound_s', outbound), ('return_s', inward), ('travel_energy_kwh', travel),
                                ('cruise_altitude_m', height), ('ground_m', ground), ('agl_m', agl)]:
            if field in station:
                audit.close(station[field], expected, 'relay_station_' + field, sid)
        audit.check(sortie['drone'] in fleet, sid + ': relay outside aircraft inventory')
        audit.check(sortie['component'] in {f'RE{k + 1:02}' for k in range(component_count)},
                    sid + ': relay energy component outside inventory')
        backhaul = geo.link(point, geo.gateway, 'backhaul')
        audit.check(backhaul['available'], sid + ': unavailable relay backhaul')
        audit.check(sortie['station'] not in relay_map, sid + ': duplicate station index unsupported by schedule schema')
        relay_map[sortie['station']] = dict(sortie=sortie, point=point, backhaul=backhaul)
        charge = recharge_seconds(energy, relay['capacity'], full_charge)
        charged = padded_return + quantized(charge)
        ready = sortie.get('component_ready_s', sortie.get('battery_ready_s', charged))
        audit.check(ready >= charged - TOL, sid + ': relay component reused before full recharge')
        if 'charge_s' in sortie:
            audit.close(sortie['charge_s'], charge, 'relay_charge_s', sid)
        if 'charge_end_s' in sortie:
            audit.check(sortie['charge_end_s'] >= charged - TOL, sid + ': relay charge_end too early')
            ready = max(ready, sortie['charge_end_s'])
        audit.grid(ready, sid + '.component_ready_s')
        air.append((sortie['drone'], sortie['start_s'], sortie['resource_end_s'], sid))
        components.append((sortie['component'], sortie['start_s'], ready, sid))
        recomputed.append(dict(id=sid, station=sortie['station'], ground_m=ground, agl_m=agl,
            outbound_s=outbound, return_duration_s=inward, return_s=padded_return, physical_return_s=physical_return,
            quantized_return_s=padded_return, return_padding_s=padded_return - physical_return,
            ground_wait_before_departure_s=max(0., wait), service_s=duration, energy_kwh=energy,
            soc_percent=soc, recharge_s=charge, earliest_component_ready_s=charged))
    audit.check(len({p['drone'] for p in data['relay']}) <= 2, 'More than two relay aircraft used')
    audit.check(len({p['component'] for p in data['relay']}) <= 6, 'More than six relay components used')
    return dict(stations=relay_map, air=air, components=components, recomputed=recomputed)


def validate_communication_senior(data, geo, physical, relay_map, audit):
    proofs = []
    minimum = math.inf
    sampled_minimum = math.inf
    sampled_witness = None
    sample_count = 0
    duration_sum = 0.
    relay_sum = 0.
    unproven = 0
    for sortie in data['transport']:
        sid = sortie['id']
        if sid not in physical:
            continue
        expected = physical[sid]['stages']
        previous = sortie['flight_start_s']
        audit.check(bool(sortie['communication']), sid + ': missing communication profile')
        used_stations = set()
        for index, row in enumerate(sortie['communication']):
            where = f'{sid}/communication[{index}]'
            audit.close(row['start'], previous, 'communication_gap_s', where)
            audit.check(row['end'] >= row['start'] - TOL, where + ': negative interval')
            stages = [s for s in expected if s['kind'] == normalized_kind(row['kind'])
                and s['leg'] == row['leg'] and s['start'] <= row['start'] + TOL and s['end'] >= row['end'] - TOL]
            audit.check(len(stages) == 1, where + ': interval does not match one independently recomputed stage')
            if stages:
                stage = stages[0]
                span = stage['end'] - stage['start']
                for field, instant in (('a', row['start']), ('b', row['end'])):
                    ratio = (instant - stage['start']) / span if span > 1e-12 else 0.
                    location = [x + ratio * (y - x) for x, y in zip(stage['a'], stage['b'])]
                    audit.close(math.dist(location, row[field]), 0., 'communication_xyz_m', where)
            station_id = row.get('station', sortie.get('station'))
            anchors = []
            reference = None
            if row['relay']:
                audit.check(station_id in relay_map, where + ': assigned relay station not scheduled')
                if station_id in relay_map:
                    used_stations.add(station_id)
                    reference = relay_map[station_id]
                    relay_sortie = reference['sortie']
                    anchors = [reference['point']]
                    begin, end = sortie['start_s'] + row['start'], sortie['start_s'] + row['end']
                    audit.check(begin >= relay_sortie['service_start_s'] - TOL and end <= relay_sortie['service_end_s'] + TOL,
                                where + ': actual relay service window does not cover complete interval')
                    minimum = min(minimum, reference['backhaul']['margin_db'])
                    relay_sum += row['end'] - row['start']
            certificate = geo.segment_certificate(row['a'], row['b'], anchors)
            audit.check(certificate['proved'], where + ': no continuous geometry certificate')
            unproven += certificate['unproven_count']
            margins = [s['min_margin_db'] for s in certificate['segments'] if 'min_margin_db' in s]
            if margins:
                minimum = min(minimum, min(margins))
            for fraction in (0., .25, .5, .75, 1.):
                point = tuple(a + fraction * (b - a) for a, b in zip(row['a'], row['b']))
                direct = geo.link(point, geo.gateway, 'direct')
                sample_count += 1
                if direct['available']:
                    margin, source = direct['margin_db'], 'G01'
                elif anchors:
                    access = geo.link(point, anchors[0], 'access')
                    backhaul = reference['backhaul']
                    margin = min(access['margin_db'], backhaul['margin_db'])
                    source = reference['sortie']['id']
                    audit.check(access['available'] and backhaul['available'], where + ': sampled relay outage')
                else:
                    margin, source = direct['margin_db'], 'outage'
                    audit.check(False, where + ': sampled direct outage')
                if margin < sampled_minimum:
                    sampled_minimum = margin
                    sampled_witness = dict(sortie=sid, interval=index, source=source, position=list(point),
                        time_s=sortie['start_s'] + row['start'] + fraction * (row['end'] - row['start']))
            if 'min_margin_db' in row:
                audit.check(row['min_margin_db'] >= -TOL, where + ': exported negative margin')
            proofs.append(dict(sortie=sid, interval=index, absolute_start_s=sortie['start_s'] + row['start'],
                absolute_end_s=sortie['start_s'] + row['end'], relay=row['relay'], station=station_id,
                certificate=certificate))
            duration_sum += row['end'] - row['start']
            previous = row['end']
        allowed_stations=data.get('communication_policy',{}).get('max_relay_stations_per_sortie',1)
        audit.check(isinstance(allowed_stations,int) and allowed_stations>=1,
                    sid + ': invalid declared relay-station search limit')
        audit.check(len(used_stations) <= allowed_stations,
                    sid + ': more relay stations than the declared search limit')
        audit.close(previous, sortie['duration_s'], 'communication_end_s', sid)
    return dict(communication_proofs=proofs, communication_interval_count=len(proofs),
        unproven_communication_interval_count=unproven, minimum_proved_communication_margin_db=minimum if math.isfinite(minimum) else None,
        minimum_sampled_effective_margin_db=sampled_minimum if math.isfinite(sampled_minimum) else None,
        minimum_sampled_margin_witness=sampled_witness, sampled_communication_points=sample_count,
        continuous_communication_aircraft_seconds=duration_sum, relay_guaranteed_aircraft_seconds=relay_sum)


def validate(solution):
    solution = Path(solution)
    audit = Audit()
    try:
        raw_solution = solution.read_bytes()
        data = json.loads(raw_solution.decode('utf-8-sig'))
    except (OSError, ValueError) as error:
        audit.check(False, 'Cannot read solution JSON: ' + str(error))
        return audit.report(solution, metrics={})
    if not schema_check(data, audit):
        return audit.report(solution, metrics={})
    audit.check(data['feasible'] is True, 'Solution does not report a feasible incumbent')
    for field, expected in EXPECTED_PHYSICS.items():
        audit.close(data['physics'][field], expected, 'physics_' + field, 'physics', 1e-10)
    geo = Geometry()
    originals = read_originals_senior()
    models, fleet, batteries, boxes = originals[:4]
    audit.check(len(boxes) == 80, 'Original box inventory is not 80')
    audit.check(Counter(fleet.values()) == Counter(A=4, B=2, C=2), 'Original aircraft inventory mismatch')
    audit.check({m: b['count'] for m, b in batteries.items()} == dict(A=6, B=4, C=4), 'Original battery inventory mismatch')
    audit.check(Counter(b for s in data['transport'] for b in s['boxes']) == Counter(boxes.keys()),
                'Boxes not assigned exactly once / unknown box')
    for name in ('transport', 'relay'):
        audit.check(len({s['id'] for s in data[name]}) == len(data[name]), name + ': duplicate sortie ids')
    for relpath, digest in data.get('inputs', {}).items():
        path = ROOT / str(relpath).replace('\\', '/')
        try:
            actual = hashlib.sha256(path.read_bytes()).hexdigest()
            audit.check(actual == digest, 'Input SHA256 mismatch: ' + str(path))
        except OSError as error:
            audit.check(False, 'Cannot verify input SHA256: ' + str(error))
    source = data['senior_source']
    if isinstance(source, dict) and 'path' in source and 'sha256' in source:
        try:
            source_path=(Path(os.environ['Q3_SENIOR_ROOT'])/'建模/outputs/q2_improved/tables/final_result.json'
                         if os.environ.get('Q3_SENIOR_ROOT') else Path(source['path']))
            actual = hashlib.sha256(source_path.read_bytes()).hexdigest()
            audit.check(actual == source['sha256'], 'Senior source SHA256 mismatch: ' + source['path'])
        except OSError as error:
            audit.check(False, 'Cannot verify senior source SHA256: ' + str(error))
    transport = recompute_transport_senior(data, geo, originals, audit)
    relay = recompute_relay_senior(data, geo, originals, audit)
    communication = validate_communication_senior(data, geo, transport['physical'], relay['stations'], audit)
    exported_deliveries = {r['box']: r for r in data['deliveries']}
    audit.check(len(exported_deliveries) == len(data['deliveries']) == len(boxes), 'Exported delivery list not unique or complete')
    audit.check(set(exported_deliveries) == set(boxes), 'Exported delivery box set mismatch')
    for bid, actual in transport['deliveries'].items():
        if not audit.check(bid in exported_deliveries, 'Missing exported delivery: ' + bid):
            continue
        row = exported_deliveries[bid]
        audit.check(row['sortie'] == actual['sortie'] and row['service'] == actual['service'], 'Delivery identity mismatch: ' + bid)
        audit.close(row['time_s'], actual['time_s'], 'exported_delivery_s', bid)
        if 'rank' in row:
            audit.check(row['rank'] == actual['rank'], 'Delivery rank mismatch: ' + bid)
        for key, source in [('due_s', 'due'), ('first_deadline_s', 'first_deadline'), ('priority', 'priority')]:
            if key in row:
                audit.check(row[key] == boxes[bid][source], 'Exported ' + key + ' differs from workbook: ' + bid)
    physical = transport['physical']
    deliveries = transport['deliveries']
    transport_energy = math.fsum(p['energy_kwh'] for p in physical.values())
    relay_energy = math.fsum(p['energy_kwh'] for p in relay['recomputed'])
    returns = [p['return_s'] for p in physical.values()] + [p['return_s'] for p in relay['recomputed']]
    metrics = dict(boxes=len(deliveries), transport_sorties=len(data['transport']), relay_sorties=len(data['relay']),
        transport_energy_kwh=transport_energy, relay_energy_kwh=relay_energy, total_energy_kwh=transport_energy + relay_energy,
        joint_finish_s=max(returns, default=0.), last_delivery_s=max((d['time_s'] for d in deliveries.values()), default=0.),
        late_boxes=sum(d['time_s'] > boxes[b]['due'] + TOL for b, d in deliveries.items()),
        weighted_delay_s=sum(max(0., d['time_s'] - boxes[b]['due']) * boxes[b]['priority'] for b, d in deliveries.items()),
        multi_stop_sorties=sum(len(s['order']) > 1 for s in data['transport']))
    for field, value in metrics.items():
        if audit.check(field in data['metrics'], 'metrics: missing ' + field):
            audit.close(data['metrics'][field], value, 'metric_' + field, 'summary')
    extra_metrics = dict(transport_finish_s=max((p['return_s'] for p in physical.values()), default=0.),
        weighted_mean_delivery_s=sum(boxes[b]['priority'] * d['time_s'] for b, d in deliveries.items())
            / sum(b['priority'] for b in boxes.values()))
    for field, value in extra_metrics.items():
        if field in data['metrics']:
            audit.close(data['metrics'][field], value, 'metric_' + field, 'summary')
        metrics[field] = value
    physical_returns = [p['return_s'] - sum(leg['quantization_wait_s'] for leg in p['legs'][-1:])
                        for p in physical.values()] + [r['physical_return_s'] for r in relay['recomputed']]
    physical_finish = max(physical_returns, default=0.)
    usage = dict(transport_aircraft=audit.overlaps(transport['air'], 'transport aircraft'),
        transport_batteries=audit.overlaps(transport['batteries'], 'transport battery including recharge'),
        relay_aircraft=audit.overlaps(relay['air'], 'relay aircraft including turnaround'),
        relay_components=audit.overlaps(relay['components'], 'relay component including recharge'))
    return audit.report(solution, solution_sha256=hashlib.sha256(raw_solution).hexdigest(),
        validator_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        metrics=metrics, senior_physics=data['physics'], senior_source=data['senior_source'],
        minimum_hard_deadline_slack_s=min((r['slack_s'] for r in transport['hard_limits']), default=None),
        hard_deadline_checks=transport['hard_limits'],
        minimum_transport_return_soc_percent=min((p['soc_percent'] for p in physical.values()), default=None),
        minimum_relay_return_soc_percent=min((p['soc_percent'] for p in relay['recomputed']), default=None),
        numeric_gaps=dict(per_sortie=transport['gaps'], total_flight_quantization_wait_s=sum(r['wait_s'] for r in transport['waits']),
            physical_joint_finish_s=physical_finish, quantized_joint_finish_s=metrics['joint_finish_s'],
            joint_finish_padding_s=metrics['joint_finish_s'] - physical_finish,
            explanation='Flight-leg millisecond ceilings create stationary numerical padding; inherited simplified energy excludes this padding.'),
        quantization_waits=transport['waits'], resource_usage=usage, recomputed_transport=physical,
        recomputed_relay=relay['recomputed'], **communication,
        assumptions=['Original workbooks reread; no senior RouteFactory, Q2 solution calculation, or Q3 solver reused.',
            'DEM supercover, radio link and continuous certificate primitives are shared with the separate Geometry module.',
            'g=9.80665 m/s^2; Euse*d/L(q) plus m*g*h/eta is the inherited simplified energy model, not physical calibration.',
            'Each full flight leg and charge duration is independently ceiled to 1 ms; per-box deliveries use actual within-stop ranks.',
            'Before-service scheduling slack is realized by waiting on the ground before departure.',
            'Relay setup consumes hover-plus-communication power for 30 s; millisecond padding has no extra modeled energy.',
            'relay=true denotes continuously guaranteed backup: direct G01 service takes priority at every instant.',
            'Reported joint completion uses millisecond-ceiled return times, excluding recharge and turnaround; physical completion is separately reported with its sub-millisecond gap.'])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--solution', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    try:
        result = validate(args.solution)
    except (KeyError, TypeError, ValueError, IndexError, OSError) as error:
        result = dict(passed=False, solution=str(args.solution), check_count=0, error_count=1,
            errors=[f'Malformed field or original-input failure ({type(error).__name__}): {error}'], warnings=[], metrics={})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    keys = ('passed', 'check_count', 'error_count', 'metrics', 'minimum_hard_deadline_slack_s',
            'minimum_transport_return_soc_percent', 'minimum_relay_return_soc_percent',
            'minimum_proved_communication_margin_db', 'unproven_communication_interval_count', 'maximum_numeric_errors')
    print(json.dumps({key: result[key] for key in keys if key in result}, ensure_ascii=True))
    if not result['passed']:
        print(json.dumps(result['errors'], ensure_ascii=True))
        raise SystemExit(1)


if __name__ == '__main__':
    main()
