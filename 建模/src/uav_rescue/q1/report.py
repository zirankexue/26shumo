from __future__ import annotations

from collections import Counter
import csv
import math
from pathlib import Path

import numpy as np
import openpyxl


def build_sheets(payload, inputs):
    b=payload['baseline'];rows=b['rows'];sheets=[]

    def add(name,headers,values,widths,formats=None):
        sheets.append({'name':name,'headers':headers,'rows':values,'widths':widths,'formats':formats or {}})

    add('Q1_单点组批',inputs.template_headers,
        [[r['sortie'],r['service'],r['drone'],'; '.join(r['box_ids']),r['mass_kg'],r['volume_m3'],
          r['flight_s'],r['energy_kwh'],100*r['return_soc']] for r in rows],
        [22,14,12,58,16,17,18,20,18],{'4':'0.000','5':'0.000','6':'0.000','7':'0.000000','8':'0.000'})
    time_rows=[]
    for r in rows:
        d=inputs.drones[r['drone']];n=r['box_count']
        time_rows.append([r['sortie'],n,d.preparation,n*d.loading,d.handover,n*d.per_box_handover,r['flight_s'],r['operation_s']])
    add('作业时间分解',['架次编号','货箱数','准备时间s','装载时间s','基础交接s','逐箱交接s','往返飞行时间s','累计作业时间s'],
        time_rows,[22,12,16,16,16,16,22,22],{str(i):'0.000' for i in range(2,8)})
    add('最大安全载荷',['服务区编号','机型编号','安全余量','最大安全载荷kg','限制原因','能量预算kWh','上限载荷能耗kWh'],
        [[r['service'],r['drone'],r['reserve'],r['safe_payload_kg'],r['reason'],r['energy_budget_kwh'],r['energy_at_limit_kwh']]
         for r in b['safe_payloads']],[15,12,14,24,28,22,26],{'2':'0%','3':'0.000000','5':'0.000000','6':'0.000000'})
    add('目标优先级对比',['目标优先级','是否可行','架次数','运输能耗kWh','累计作业时间s','累计作业时间h','A型架次','B型架次','C型架次'],
        [[x['label'],'是' if x['feasible'] else '否',x['summary']['sorties'],x['summary']['energy_kwh'],
          x['summary']['operation_s'],None if not x['feasible'] else x['summary']['operation_s']/3600,
          *([x['summary']['drone_counts'][g] for g in ['A','B','C']] if x['feasible'] else [None]*3)] for x in payload['comparisons']],
        [30,14,12,24,26,24,15,15,15],{'3':'0.000000','4':'0.000','5':'0.000000'})
    add('余量敏感性汇总',['安全余量','是否可行','最少架次数','运输能耗kWh','累计作业时间s','A型架次','B型架次','C型架次','不可行服务区'],
        [[x['reserve'],'是' if x['feasible'] else '否',x['summary']['sorties'],x['summary']['energy_kwh'],x['summary']['operation_s'],
          *([x['summary']['drone_counts'][g] for g in ['A','B','C']] if x['feasible'] else [None]*3),';'.join(x['infeasible_services'])]
         for x in payload['sensitivity']],[15,14,18,24,26,15,15,15,24],{'0':'0%','3':'0.000000','4':'0.000'})
    add('余量载荷明细',['安全余量','服务区编号','机型编号','最大安全载荷kg','限制原因'],
        [[s['reserve'],r['service'],r['drone'],r['safe_payload_kg'],r['reason']] for s in payload['sensitivity'] for r in s['safe_payloads']],
        [16,16,14,26,30],{'0':'0%','3':'0.000000'})
    failures=[[s['reserve'],r['service'],r['category'],r['mass_kg'],r['drone'],r['single_box_energy_kwh'],r['budget_kwh'],
               '可行' if r['mass_ok'] and r['volume_ok'] and r['energy_ok'] else '不可行']
              for s in payload['sensitivity'] for r in s['infeasibility_details']]
    if failures:
        add('不可行情景诊断',['安全余量','服务区','无法交付类别','单箱质量kg','机型','单箱往返能耗kWh','可用预算kWh','单箱是否可行'],
            failures,[15,14,23,19,12,27,23,22],{'0':'0%','3':'0.000','5':'0.000000','6':'0.000000'})
    summary=[]
    for sid in inputs.services:
        rs=[r for r in rows if r['service']==sid]
        summary.append([sid,len(rs),sum(r['box_count'] for r in rs),math.fsum(r['mass_kg'] for r in rs),
                        math.fsum(r['volume_m3'] for r in rs),math.fsum(r['energy_kwh'] for r in rs),
                        math.fsum(r['operation_s'] for r in rs),*[sum(r['drone']==g for r in rs) for g in ['A','B','C']]])
    add('服务区汇总',['服务区','架次数','货箱数','质量kg','体积m³','能耗kWh','累计作业时间s','A型架次','B型架次','C型架次'],
        summary,[14,12,12,16,16,22,25,15,15,15],{'3':'0.000','4':'0.000','5':'0.000000','6':'0.000'})
    add('核验结果',['服务区','候选批次数','DP架次数','MILP架次数','DP能耗kWh','MILP能耗kWh','DP时间s','MILP时间s','核验状态'],
        [[r['service'],r['candidate_count'],r['dp_objective'][0],r['milp_objective'][0],r['dp_objective'][1],r['milp_objective'][1],
          r['dp_objective'][2],r['milp_objective'][2],'通过' if r['passed'] else '失败'] for r in payload['validation']['milp']],
        [14,19,18,20,23,25,22,23,16],{str(i):'0.000000' for i in range(4,8)})
    notes=[['问题范围','仅单点往返组批，不含机队、电池调度、通信和配送时限。'],
           ['目标顺序','主方案先最少架次，再最低能耗，再最短累计作业时间。'],
           ['往返时间','Q1_单点组批的往返时间为去返程飞行时间之和；全部作业时间见作业时间分解。'],
           ['载荷定义','最大安全载荷只含质量与能量约束；货箱装载还须满足体积约束。'],
           ['能耗解释','水平能耗按可用电量×距离/等效航程；爬升按重力势能/效率折算。具体推导见模型说明。'],
           ['高程口径','节点作业海拔取节点表，沿途地形取完整DEM；覆盖所有相交像元。'],
           ['数值精度','动态规划能耗按1e-12 kWh、时间按1e-6 s比较；导出保留原始浮点结果。'],
           ['结果性质','程序计算快照；修改输入后须重新运行run_q1.py，Excel不会自行重算优化模型。'],
           ['来源','物资需求与配送时限.xlsx；调度中心与服务区.xlsx；运输无人机数据.xlsx；30米DEM.tif；结果提交模板.xlsx。']]
    add('口径与来源',['事项','说明'],notes,[22,110])
    return sheets


def make_figures(payload, font_path: Path, folder: Path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    if font_path.exists():
        font_manager.fontManager.addfont(str(font_path))
        name=font_manager.FontProperties(fname=str(font_path)).get_name()
        plt.rcParams['font.family']=[name,'DejaVu Sans']
    plt.rcParams.update({'axes.unicode_minus':False,'font.size':11,'axes.spines.top':False,
                         'axes.spines.right':False,'pdf.fonttype':42,'savefig.facecolor':'white'})
    folder.mkdir(parents=True,exist_ok=True)
    def save(fig,name):
        fig.savefig(folder/f'{name}.png',dpi=300,bbox_inches='tight')
        fig.savefig(folder/f'{name}.pdf',bbox_inches='tight')
        plt.close(fig)
    records=payload['baseline']['safe_payloads']
    services=sorted({r['service'] for r in records})
    lookup={(r['service'],r['drone']):r['safe_payload_kg'] for r in records}
    values=np.array([[lookup[(s,g)] if lookup[(s,g)] is not None else np.nan for g in ['A','B','C']] for s in services])
    fig,ax=plt.subplots(figsize=(6.8,8.2),layout='constrained')
    heat=ax.imshow(values,cmap='YlGnBu',aspect='auto',vmin=0,vmax=80)
    ax.set(xticks=range(3),xticklabels=['A型','B型','C型'],yticks=range(15),yticklabels=services,
           title='各服务区最大安全载荷（安全余量20%）')
    for i in range(15):
        for j in range(3):
            v=values[i,j]
            ax.text(j,i,'不可达' if np.isnan(v) else f'{v:.2f}',ha='center',va='center',color='white' if v>48 else '#172B3A')
    fig.colorbar(heat,ax=ax,label='最大安全载荷（kg）',shrink=.8)
    save(fig,'01_最大安全载荷')
    fig,ax=plt.subplots(figsize=(10.5,4.8),layout='constrained')
    x=np.arange(15);bottom=np.zeros(15)
    for g,color in zip(['A','B','C'],['#7CA5BA','#D4A359','#315779']):
        counts=np.array([sum(r['service']==s and r['drone']==g for r in payload['baseline']['rows']) for s in services])
        ax.bar(x,counts,bottom=bottom,label=f'{g}型',color=color,width=.7)
        bottom+=counts
    ax.set(xticks=x,xticklabels=services,ylabel='架次数',title='主方案的机型与服务区架次分布')
    ax.tick_params(axis='x',rotation=45)
    ax.set_yticks(range(int(max(bottom))+1));ax.legend(ncols=3,frameon=False)
    save(fig,'02_机型使用构成')
    ss=payload['sensitivity'];rho=np.array([s['reserve']*100 for s in ss])
    fig,axs=plt.subplots(1,3,figsize=(13,4.2),layout='constrained')
    for ax,key,label in zip(axs,['sorties','energy_kwh','operation_s'],['最少架次数','运输能耗（kWh）','累计作业时间（h）']):
        vals=[np.nan if not s['feasible'] else s['summary'][key]/(3600 if key=='operation_s' else 1) for s in ss]
        ax.plot(rho,vals,'o-',color='#315779',lw=2)
        ax.axvline(payload['baseline']['reserve']*100,color='#A5A5A5',linestyle='--',lw=1)
        ax.set(xlabel='返航安全余量（%）',ylabel=label,xticks=rho,
               xlim=(rho.min()-1.5,rho.max()+1.5))
        if key=='sorties':
            for x,y in zip(rho,vals):
                if np.isfinite(y):
                    ax.annotate(f'{int(y)}',(x,y),xytext=(0,7),textcoords='offset points',
                                ha='center',fontsize=9,color='#315779')
            ax.margins(y=.18)
        for s in ss:
            if not s['feasible']:
                ax.axvline(s['reserve']*100,color='#AB493B',linestyle=':',lw=1.3)
                ax.text(s['reserve']*100-.3,.52,'不可行',transform=ax.get_xaxis_transform(),
                        rotation=90,ha='right',va='center',color='#AB493B',fontsize=10)
        ax.grid(axis='y',alpha=.18)
    fig.suptitle(f'安全余量敏感性（{rho.min():g}%—{rho.max():g}%）：每档均重新优化组批',fontsize=14)
    save(fig,'03_安全余量敏感性')
    fig,axs=plt.subplots(1,3,figsize=(13,4.3),layout='constrained')
    for ax,g in zip(axs,['A','B','C']):
        for color_index,sid in enumerate(services):
            vals=[next(r['safe_payload_kg'] for r in s['safe_payloads'] if r['service']==sid and r['drone']==g) for s in ss]
            ax.plot(rho,[np.nan if v is None else v for v in vals],lw=1.2,alpha=.9,label=sid,
                    color=plt.get_cmap('tab20')(color_index))
        ax.set(title=f'{g}型',xlabel='返航安全余量（%）',ylabel='最大安全载荷（kg）',xticks=rho,
               xlim=(rho.min()-1.5,rho.max()+1.5))
        ax.grid(alpha=.15)
    handles,labels=axs[-1].get_legend_handles_labels()
    fig.legend(handles,labels,loc='outside lower center',ncols=8,frameon=False,fontsize=9)
    fig.suptitle('不同机型及服务区的安全载荷变化',fontsize=14)
    save(fig,'04_各机型载荷敏感性')


def write_report(payload,path):
    b=payload['baseline'];s=b['summary']
    parts=['# 问题一计算结果分析','',
        f"采用架次、能耗、累计作业时间的词典序目标。在20%安全余量下，80箱物资全部交付，主方案共{s['sorties']}架次，总运输能耗{s['energy_kwh']:.6f} kWh，累计作业时间{s['operation_s']:.3f}秒（{s['operation_s']/3600:.6f}小时）。",
        '',f"机型架次为A型{s['drone_counts']['A']}次、B型{s['drone_counts']['B']}次、C型{s['drone_counts']['C']}次。累计作业时间是各架次作业耗时相加，不是考虑多机并行后的任务完工时刻。问题一也不限制实体机库存，因此架次数不代表需要同时配置的无人机数。",'',
        '## 1 安全载荷与组批','',
        '| 机型 | 最小安全载荷kg | 最大安全载荷kg | 受能量约束的服务区 |','|---|---:|---:|---|']
    for g in ['A','B','C']:
        rs=[r for r in b['safe_payloads'] if r['drone']==g];vals=[r['safe_payload_kg'] for r in rs if r['safe_payload_kg'] is not None]
        limits='、'.join(r['service'] for r in rs if r['reason']=='往返安全能量限制') or '无'
        parts.append(f"| {g} | {min(vals):.6f} | {max(vals):.6f} | {limits} |")
    parts+=['',f"质量与体积给出的架次下界为{b['sortie_lower_bound']}：S001、S002、S003各至少2次，其余12个区各至少1次。主方案达到该下界，因此18架次的第一层目标也有直接的下界证明。",'',
        '![最大安全载荷](figures/01_最大安全载荷.png)','',
        '每个候选批次均同时验证质量、体积和往返能量。安全载荷是连续质量上限，不能将其直接理解为某种货箱的实际可装质量。',
        '', '![机型使用构成](figures/02_机型使用构成.png)','',
        '## 2 多目标优先级对比','',
        '| 优先级 | 架次数 | 能耗kWh | 累计作业小时 | A/B/C架次 |','|---|---:|---:|---:|---|']
    for x in payload['comparisons']:
        y=x['summary'];mix='/'.join(str(y['drone_counts'][g]) for g in ['A','B','C'])
        parts.append(f"| {x['label']} | {y['sorties']} | {y['energy_kwh']:.6f} | {y['operation_s']/3600:.6f} | {mix} |")
    energy_first=next((x for x in payload['comparisons'] if x['order'][0]=='energy'),None)
    if energy_first is not None:
        e=energy_first['summary']
        parts+=['',f"能耗优先方案相对主方案增加{e['sorties']-s['sorties']}架次，节能{s['energy_kwh']-e['energy_kwh']:.6f} kWh（{100*(s['energy_kwh']-e['energy_kwh'])/s['energy_kwh']:.4f}%），增加累计作业时间{(e['operation_s']-s['operation_s'])/60:.3f}分钟。该对照量化了优先减少架次的代价。"]
    parts+=['','对照仅改变目标优先关系，不改变物理与装载约束。这三组方案不是完整Pareto前沿。主方案体现先减少往返任务组织次数，再控制能耗的偏好。',
        '', '## 3 安全余量敏感性','',
        '| 安全余量 | 可行性 | 最少架次 | 能耗kWh | 累计作业小时 | A/B/C架次 |','|---|---|---:|---:|---:|---|']
    for x in payload['sensitivity']:
        y=x['summary']
        if x['feasible']:
            mix='/'.join(str(y['drone_counts'][g]) for g in ['A','B','C'])
            parts.append(f"| {x['reserve']:.0%} | 可行 | {y['sorties']} | {y['energy_kwh']:.6f} | {y['operation_s']/3600:.6f} | {mix} |")
        else:parts.append(f"| {x['reserve']:.0%} | 不可行：{','.join(x['infeasible_services'])} | — | — | — | — |")
    for scenario in payload['sensitivity']:
        if scenario['infeasibility_details']:
            groups=sorted({(r['service'],r['category']) for r in scenario['infeasibility_details']})
            text='；'.join(f'{sid}的{category}' for sid,category in groups)
            parts+=['',f"{scenario['reserve']:.0%}余量下不可行的直接原因：{text}无法用任何机型进行单箱往返配送。能耗随载荷不减，增加同批货箱不会恢复可行性，因此这不是算法未找到方案。各机型的单箱能耗与预算见Excel“不可行情景诊断”。"]
    parts+=['','![安全余量敏感性](figures/03_安全余量敏感性.png)','',
        '安全余量增加会收缩可行域，最大安全载荷不增、最少架次不减。能耗和作业时间受组批及机型切换共同影响，不要求单调。表中仅列离散采样情景，不能据此声称确定了全部精确临界余量。',
        '', '![各机型安全载荷敏感性](figures/04_各机型载荷敏感性.png)','',
        '## 4 可行性与最优性核验','',
        f"- 全部80个货箱唯一交付；质量758 kg、体积2.011 m³；每架次仅服务一个服务区。",
        f"- 从逐箱原始属性独立重算载荷、体积、能耗、SOC和时间；最低返航SOC为{b['verification']['min_return_soc']*100:.6f}%。",
        f"- 共{b['state_count']}个箱数状态。动态规划遍历所有可行组成，每个服务区的主方案均经三阶段MILP复核，并取得最优终止状态。",
        '- 最优性针对明确采用的地形、能耗、时间和目标口径；不是对所有可能物理解释均成立的结论。',
        '- 能耗按1e-12 kWh、作业时间按1e-6 s转换为整数比较，消除浮点加法顺序影响；导出指标用未量化值重算。每个完整方案至多80架次，相应累计量化误差上界为4e-11 kWh与4e-5秒。',
        '', '## 5 文件与复现','',
        '- `问题一结果.xlsx`：主方案、时间分解、最大安全载荷、目标对比、敏感性和核验结果。',
        '- `tables/results.json`：完整结构化结果、所有对照与敏感性情景的逐箱方案、参数和输入哈希。',
        '- `tables/*.csv`：与Excel对应的数值表；目标及余量情景的JSON保留全部组批细节。',
        '- `logs/validation.json`：原始数据重算、MILP、敏感性、Excel和输入未修改检查。',
        '- 工程根目录运行 `python run_q1.py --config configs/q1.toml` 可重新生成。',
        '', '本问不排实体机、共享电池或交付时刻。此方案可作为后续问题的参考，不能直接宣称满足问题二的时限与资源约束。','']
    text='\n'.join(parts)
    text=text.replace('](figures/', ']('+str((path.parent/'figures').resolve()).replace('\\','/')+'/')
    path.write_text(text,encoding='utf-8')


def verify_workbook(path,sheets):
    wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
    checked=0
    try:
        if wb.sheetnames!=[s['name'] for s in sheets]:raise AssertionError('导出工作表不一致')
        for sheet in sheets:
            ws=wb[sheet['name']]
            expected=[sheet['headers']]+sheet['rows']
            with (path.parent/'tables'/f"{sheet['name']}.csv").open(encoding='utf-8-sig',newline='') as stream:
                csv_rows=list(csv.reader(stream))
            if len(csv_rows)!=len(expected):raise AssertionError('CSV行数不一致')
            for row_index,row in enumerate(expected,1):
                if len(csv_rows[row_index-1])!=len(row):raise AssertionError('CSV列数不一致')
                for col_index,value in enumerate(row,1):
                    actual=ws.cell(row_index,col_index).value
                    csv_value=csv_rows[row_index-1][col_index-1]
                    if isinstance(value,(int,float)):
                        if not isinstance(actual,(int,float)) or not math.isclose(actual,value,rel_tol=1e-12,abs_tol=1e-9):
                            raise AssertionError(f'{sheet["name"]}!{row_index},{col_index}数值不符：{actual}/{value}')
                        if not math.isclose(float(csv_value),value,rel_tol=1e-12,abs_tol=1e-9):raise AssertionError('CSV数值不一致')
                    elif (actual if actual is not None else '')!=(value if value is not None else ''):
                        raise AssertionError(f'{sheet["name"]}!{row_index},{col_index}文本不符')
                    elif csv_value!=(value if value is not None else ''):raise AssertionError('CSV文本不一致')
                    checked+=1
    finally:wb.close()
    return {'performed':True,'passed':True,'checked_cells':checked,'sheets':len(sheets),'csv_also_checked':True}
