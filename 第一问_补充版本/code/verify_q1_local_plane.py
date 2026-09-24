"""Independent original-input audit for the Q1 local ellipsoidal plane revision.

First run --geometry-only after the baseline solver, then run with no arguments
after the multiobjective and reserve-sensitivity solvers. No solver is imported.
"""
from __future__ import annotations

import argparse
import hashlib
import itertools
import json
import math
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, getcontext
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if (ROOT / '.runtime_py').is_dir():
    sys.path.insert(0, str(ROOT / '.runtime_py'))
import numpy as np
import openpyxl
from PIL import Image

getcontext().prec = 40
D = lambda value: Decimal(str(value))
TYPES = ('MED', 'WAT', 'FOD', 'HYG')
ENERGY_TOL, TIME_TOL = 1e-9, 1e-6


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def read_originals(solution):
    paths = {key: ROOT / value['path'] for key, value in solution['sources'].items()}
    assert all(digest(paths[key]) == value['sha256'] for key, value in solution['sources'].items())
    nodes = {}
    with_workbook = openpyxl.load_workbook(paths['nodes'], read_only=True, data_only=True)
    for row in with_workbook['数据'].values:
        if row[0] == 'O01' or (isinstance(row[0], str) and row[0].startswith('S0')):
            nodes[row[0]] = dict(id=row[0], lon=float(row[2]), lat=float(row[3]), elevation_m=float(row[4]))
    with_workbook.close()
    fields = ('id', 'name', 'empty_mass_kg', 'payload_kg', 'volume_m3', 'cruise_mps',
              'empty_range_m', 'full_range_m', 'energy_kwh', 'reserve_percent',
              'prepare_s', 'load_per_box_s', 'handover_base_s', 'handover_per_box_s',
              'climb_mps', 'descend_mps', 'climb_efficiency', 'descent_extra_efficiency')
    wb = openpyxl.load_workbook(paths['drones'], read_only=True, data_only=True)
    drones = {row[0]: dict(zip(fields, row)) for row in wb['数据'].iter_rows(min_row=3, max_row=5, values_only=True)}
    wb.close()
    wb = openpyxl.load_workbook(paths['boxes'], read_only=True, data_only=True)
    boxes = [dict(id=row[0], service_id=row[1], type=row[0].split('-')[1],
                  mass_kg=float(row[3]), volume_m3=float(row[4]))
             for row in wb['逐箱货箱清单'].iter_rows(min_row=2, values_only=True) if row[0]]
    wb.close()
    assert len(boxes) == len({box['id'] for box in boxes}) == 80
    return paths, nodes, drones, boxes


def closed_rectangle_intersects(x0, y0, x1, y1, col, row):
    """Slab clipping against a closed unit rectangle; independent of supercover."""
    lower, upper = 0.0, 1.0
    for origin, delta, bound in ((x0, x1-x0, col), (y0, y1-y0, row)):
        if delta == 0:
            if not bound <= origin <= bound + 1:
                return False
        else:
            t0, t1 = (bound-origin)/delta, (bound+1-origin)/delta
            lower, upper = max(lower, min(t0, t1)), min(upper, max(t0, t1))
            if lower > upper + 1e-13:
                return False
    return True


def geometry_audit(solution_path, solution, paths, nodes):
    with Image.open(paths['dem']) as image:
        dem = np.array(image)
        scale, tie, keys = image.tag_v2[33550], image.tag_v2[33922], image.tag_v2[34735]
    entries = {keys[i]: keys[i+3] for i in range(4, len(keys), 4)}
    assert entries[2048] == 4326 and entries[1025] == 2
    lon_center, lat_center = tie[3]-tie[0]*scale[0], tie[4]+tie[1]*scale[1]
    origin = nodes['O01']
    phi = math.radians(origin['lat'])
    a, e2 = 6378137.0, 6.6943799901413165e-3
    denominator = math.sqrt(1-e2*math.sin(phi)**2)
    east_scale = a / denominator * math.cos(phi)
    north_scale = a * (1-e2) / denominator**3
    grid = lambda node: ((node['lon']-lon_center)/scale[0]+0.5,
                         (lat_center-node['lat'])/scale[1]+0.5)
    x0, y0 = grid(origin)
    independent, records = {}, {}
    for sid, node in sorted(nodes.items()):
        if sid == 'O01':
            continue
        x1, y1 = grid(node)
        cells = set()
        for row in range(max(0, math.floor(min(y0, y1))-1), min(dem.shape[0], math.floor(max(y0, y1))+2)):
            for col in range(max(0, math.floor(min(x0, x1))-1), min(dem.shape[1], math.floor(max(x0, x1))+2)):
                if closed_rectangle_intersects(x0, y0, x1, y1, col, row):
                    cells.add((row, col))
        heights = [float(dem[row, col]) for row, col in cells]
        assert heights and all(math.isfinite(h) and h != -32767 for h in heights)
        peak = max(heights)
        peak_cells = sorted([list(cell) for cell in cells if float(dem[cell]) == peak])
        values = dict(distance_m=math.hypot(east_scale*math.radians(node['lon']-origin['lon']),
                                          north_scale*math.radians(node['lat']-origin['lat'])),
                      max_terrain_m=peak, cruise_altitude_m=peak+50,
                      outbound_climb_m=peak+50-origin['elevation_m'],
                      return_climb_m=peak+50-node['elevation_m']-30)
        assert values['outbound_climb_m'] >= 0 and values['return_climb_m'] >= 0
        stored = solution['routes'][sid]
        stored_cells = set(map(tuple, stored['crossed_cells']))
        differences = {key: abs(value-stored[key]) for key, value in values.items()}
        independent[sid] = values
        records[sid] = dict(independent_values=values, absolute_differences=differences,
                            independent_cell_count=len(cells), solver_cell_count=len(stored_cells),
                            cell_sets_equal=cells == stored_cells,
                            only_in_independent=sorted(cells-stored_cells), only_in_solver=sorted(stored_cells-cells),
                            independent_peak_cells=peak_cells, peak_cells_equal=peak_cells == sorted(stored['peak_cells']))
    report = dict(audit_created_utc=datetime.now(timezone.utc).isoformat(),
                  auditor='verify_q1_local_plane.py',
                  method='Re-read original XLSX/GeoTIFF. WGS84 fixed N(phi0)cos(phi0), M(phi0) at O01. Enumerate every bounding-box pixel and independently clip the straight segment against each closed pixel rectangle; no solver geometry functions or stored geometry inputs are used.',
                  pixel_registration='PixelIsPoint tiepoint denotes pixel center; add 0.5 before closed-cell indexing.',
                  sources=dict(solution_sha256=digest(solution_path), nodes_sha256=digest(paths['nodes']), dem_sha256=digest(paths['dem'])),
                  local_plane=dict(origin='O01', origin_lon=origin['lon'], origin_lat=origin['lat'],
                                   semi_major_axis_m=a, eccentricity_squared=e2,
                                   east_scale_m_per_radian=east_scale, north_scale_m_per_radian=north_scale),
                  summary=dict(all_cell_sets_equal=all(r['cell_sets_equal'] for r in records.values()),
                               all_peak_cells_equal=all(r['peak_cells_equal'] for r in records.values()), routes_audited=len(records),
                               max_abs_difference_by_field={key:max(r['absolute_differences'][key] for r in records.values()) for key in values},
                               independent_total_route_cell_visits=sum(r['independent_cell_count'] for r in records.values())), routes=records)
    assert report['summary']['all_cell_sets_equal'] and report['summary']['all_peak_cells_equal']
    assert max(report['summary']['max_abs_difference_by_field'].values()) < 1e-8
    write_json(solution_path.parent/'geometry_audit.json', report)
    return independent, report


class IndependentPhysics:
    def __init__(self, drones, routes):
        self.drones, self.routes = drones, routes

    @lru_cache(maxsize=None)
    def cost(self, sid, model, mass, count):
        d = {key: D(value) for key, value in self.drones[model].items() if isinstance(value, (float, int))}
        r = {key: D(value) for key, value in self.routes[sid].items()}
        ratio = D(mass)/d['payload_kg']
        equivalent_range = d['empty_range_m']-(d['empty_range_m']-d['full_range_m'])*ratio*ratio.sqrt()
        energy = d['energy_kwh']*r['distance_m']*(1/equivalent_range+1/d['empty_range_m'])
        energy += D('9.80665')*((d['empty_mass_kg']+D(mass))*r['outbound_climb_m']+d['empty_mass_kg']*r['return_climb_m'])/(d['climb_efficiency']*D(3600000))
        flight = 2*r['distance_m']/d['cruise_mps']+(r['outbound_climb_m']+r['return_climb_m'])*(1/d['climb_mps']+1/d['descend_mps'])
        work = flight+d['prepare_s']+d['handover_base_s']+D(count)*(d['load_per_box_s']+d['handover_per_box_s'])
        reserve = 100*(1-energy/d['energy_kwh'])
        return energy, flight, work, reserve

    def safe_payload(self, sid, model, rho):
        drone = self.drones[model]
        budget = D(drone['energy_kwh'])*(1-D(rho)/100)
        upper = D(drone['payload_kg'])
        if self.cost(sid, model, Decimal(0), 0)[0] > budget:
            return None, 'empty_roundtrip_infeasible'
        if self.cost(sid, model, upper, 0)[0] <= budget:
            return float(upper), 'rated_payload'
        lower = Decimal(0)
        while upper-lower > D('1e-10'):
            midpoint = (lower+upper)/2
            if self.cost(sid, model, midpoint, 0)[0] <= budget:
                lower = midpoint
            else:
                upper = midpoint
        return float(lower), 'energy_reserve'


def nondominated(values):
    front = []
    for value in sorted(values):
        if any(old[0] <= value[0] and old[1] <= value[1]+ENERGY_TOL and old[2] <= value[2]+TIME_TOL for old in front):
            continue
        front = [old for old in front if not (value[0] <= old[0] and value[1] <= old[1]+ENERGY_TOL and value[2] <= old[2]+TIME_TOL)]
        front.append(value)
    return front


def bitmask_frontier(sid, boxes, drones, physics):
    """Each bit is one real box ID; least-significant-bit anchoring avoids permutations."""
    size = 1 << len(boxes)
    options = [[] for _ in range(size)]
    masses, volumes = [Decimal(0)]*size, [Decimal(0)]*size
    for mask in range(1, size):
        bit = mask & -mask
        box = boxes[bit.bit_length()-1]
        masses[mask] = masses[mask^bit]+D(box['mass_kg'])
        volumes[mask] = volumes[mask^bit]+D(box['volume_m3'])
        for model, drone in drones.items():
            if masses[mask] > D(drone['payload_kg']) or volumes[mask] > D(drone['volume_m3']):
                continue
            energy, _, work, _ = physics.cost(sid, model, masses[mask], mask.bit_count())
            if energy <= D(drone['energy_kwh'])*(1-D(drone['reserve_percent'])/100)+D('1e-10'):
                options[mask].append((1, float(energy), float(work)))
        options[mask] = nondominated(options[mask])
    fronts = [[] for _ in range(size)]
    fronts[0] = [(0, 0.0, 0.0)]
    transitions = 0
    for mask in range(1, size):
        bit = mask & -mask
        rest, sub = mask ^ bit, mask ^ bit
        front = []
        while True:
            chosen = sub | bit
            remaining = mask ^ chosen
            for _, energy, work in options[chosen]:
                for n, previous_energy, previous_work in fronts[remaining]:
                    transitions += 1
                    candidate = (n+1, energy+previous_energy, work+previous_work)
                    if not any(old[0] <= candidate[0] and old[1] <= candidate[1]+ENERGY_TOL and old[2] <= candidate[2]+TIME_TOL for old in front):
                        front = [old for old in front if not (candidate[0] <= old[0] and candidate[1] <= old[1]+ENERGY_TOL and candidate[2] <= old[2]+TIME_TOL)]
                        front.append(candidate)
            if not sub:
                break
            sub = (sub-1) & rest
        fronts[mask] = front
    return sorted(fronts[-1]), dict(service_id=sid, real_box_ids=[b['id'] for b in boxes], bitmask_state_count=size,
                                  feasible_masks=sum(bool(x) for x in options), label_transitions=transitions)


def assert_objective(actual, expected):
    assert actual[0] == expected[0], (actual, expected)
    assert abs(actual[1]-expected[1]) < 1e-8, (actual, expected)
    assert abs(actual[2]-expected[2]) < 1e-5, (actual, expected)


def objective(record):
    return record['flight_count'], record['energy_kwh'], record['work_time_s']


def verify_plan(plan, boxes, drones, physics):
    by_id = {box['id']:box for box in boxes}
    ids = [bid for batch in plan['batches'] for bid in batch['box_ids']]
    assert Counter(ids) == Counter(by_id.keys())
    differences = dict(energy_kwh=0.0, flight_time_s=0.0, work_time_s=0.0, return_soc_percent=0.0)
    sums = defaultdict(Decimal)
    soc = []
    for batch in plan['batches']:
        cargo = [by_id[bid] for bid in batch['box_ids']]
        assert all(box['service_id'] == batch['service_id'] for box in cargo)
        mass, volume = sum(D(b['mass_kg']) for b in cargo), sum(D(b['volume_m3']) for b in cargo)
        drone = drones[batch['model']]
        assert mass <= D(drone['payload_kg']) and volume <= D(drone['volume_m3'])
        assert abs(float(mass)-batch['mass_kg']) < 1e-10 and abs(float(volume)-batch['volume_m3']) < 1e-12
        cost = physics.cost(batch['service_id'], batch['model'], mass, len(cargo))
        for key, value in zip(differences, cost):
            differences[key] = max(differences[key], abs(float(value)-batch[key]))
            if key != 'return_soc_percent':
                sums[key] += value
        rho = batch.get('reserve_percent', plan.get('reserve_percent', drone['reserve_percent']))
        assert cost[0] <= D(drone['energy_kwh'])*(1-D(rho)/100)+D('1e-9')
        sums['mass_kg'] += mass
        sums['volume_m3'] += volume
        soc.append(cost[-1])
    totals = dict(flight_count=len(plan['batches']), box_count=len(ids),
                  **{key:float(value) for key,value in sums.items()},
                  minimum_soc_percent=float(min(soc)), model_counts=dict(Counter(b['model'] for b in plan['batches'])))
    stored = plan.get('totals', plan)
    for key in ('energy_kwh', 'flight_time_s', 'work_time_s', 'mass_kg', 'volume_m3', 'minimum_soc_percent'):
        if key in stored:
            assert abs(stored[key]-totals[key]) < (1e-5 if 'time' in key else 1e-8), (key, stored[key], totals[key])
    assert stored['flight_count'] == totals['flight_count']
    assert max(differences.values()) < 1e-6
    return dict(passed=True, batch_count=len(plan['batches']), recomputed_totals=totals, maximum_per_batch_difference=differences)


class QuantityAudit:
    """Independent Decimal candidate costs, quantity states used for exhaustive rho events."""
    def __init__(self, boxes, drones, physics):
        self.pools, self.demands, self.states, self.cache = {}, {}, {}, {}
        attributes = {}
        for typ in TYPES:
            values = {(D(b['mass_kg']), D(b['volume_m3'])) for b in boxes if b['type'] == typ}
            assert len(values) == 1
            attributes[typ] = values.pop()
        for sid in sorted(physics.routes):
            demand = tuple(sum(b['service_id'] == sid and b['type'] == typ for b in boxes) for typ in TYPES)
            self.demands[sid] = demand
            self.states[sid] = sorted(itertools.product(*(range(n+1) for n in demand)), key=lambda s:(sum(s), s))
            pool = []
            for counts in self.states[sid][1:]:
                mass = sum(attributes[typ][0]*n for typ,n in zip(TYPES, counts))
                volume = sum(attributes[typ][1]*n for typ,n in zip(TYPES, counts))
                for model, drone in drones.items():
                    if mass > D(drone['payload_kg']) or volume > D(drone['volume_m3']):
                        continue
                    energy, _, work, rho = physics.cost(sid, model, mass, sum(counts))
                    pool.append(dict(counts=counts, model=model, energy=energy, time=work, rho=rho))
            self.pools[sid] = pool

    def local(self, sid, rho):
        pool = self.pools[sid]
        feasible = tuple(i for i,p in enumerate(pool) if p['rho']+D('1e-10') >= D(rho))
        key = sid, feasible
        if key not in self.cache:
            dp = {(0,0,0,0):(0,Decimal(0),Decimal(0))}
            for state in self.states[sid][1:]:
                best = None
                for index in feasible:
                    p = pool[index]
                    remain = tuple(x-y for x,y in zip(state,p['counts']))
                    old = dp.get(remain)
                    if old is None:
                        continue
                    new = old[0]+1, old[1]+p['energy'], old[2]+p['time']
                    if best is None or new[0] < best[0] or (new[0] == best[0] and
                        (new[1] < best[1]-D('1e-10') or (abs(new[1]-best[1]) <= D('1e-10') and new[2] < best[2]-D('1e-7')))):
                        best = new
                if best is not None:
                    dp[state] = best
            self.cache[key] = dp.get(self.demands[sid])
        return self.cache[key]

    def global_objective(self, rho):
        local = [self.local(sid,rho) for sid in self.pools]
        if any(item is None for item in local):
            return None
        return sum(item[0] for item in local), float(sum(item[1] for item in local)), float(sum(item[2] for item in local))


def full_audit(solution_path, solution, paths, drones, boxes, routes, geometry):
    started = time.perf_counter()
    physics = IndependentPhysics(drones, routes)
    baseline_check = verify_plan(solution, boxes, drones, physics)
    local_frontiers, statistics = {}, []
    for sid in routes:
        local_frontiers[sid], stat = bitmask_frontier(sid, [b for b in boxes if b['service_id'] == sid], drones, physics)
        statistics.append(stat)
        print('BITMASK', sid, stat['bitmask_state_count'], len(local_frontiers[sid]), flush=True)
    global_frontier = [(0,0.0,0.0)]
    for local in local_frontiers.values():
        global_frontier = nondominated([(a[0]+b[0],a[1]+b[1],a[2]+b[2]) for a in global_frontier for b in local])
    global_frontier.sort()
    assert_objective(min(global_frontier), objective(solution['totals']))
    baseline_payloads = []
    for summary in solution['service_summary']:
        for model, stored in summary['safe_payloads'].items():
            value, label = physics.safe_payload(summary['service_id'], model, drones[model]['reserve_percent'])
            assert label == stored['limiting_factor']
            assert (value is None and stored['safe_payload_kg'] is None) or abs(value-stored['safe_payload_kg']) <= 1.1e-6
            baseline_payloads.append(dict(service_id=summary['service_id'],model=model,safe_payload_kg=value,limiting_factor=label))
    shared = dict(created_utc=datetime.now(timezone.utc).isoformat(), passed=True, issues=[],
                  source_sha256={key:digest(path) for key,path in paths.items()},
                  solution_sha256=digest(solution_path), geometry_audit_sha256=digest(solution_path.parent/'geometry_audit.json'),
                  decimal_precision_digits=40, gravity_m_s2=9.80665)
    baseline_audit = dict(shared, audit_type='Independent original-workbook, exact box-ID bitmask Pareto DP and Decimal physics; independently verified local-plane geometry.',
                          **baseline_check, independent_global_optimum=global_frontier[0],
                          independent_global_frontier=global_frontier, services=statistics, safe_payloads=baseline_payloads)
    write_json(solution_path.parent/'independent_audit.json', baseline_audit)
    multi_path = ROOT/'results/question1_multiobjective/multiobjective.json'
    multi = json.loads(multi_path.read_text(encoding='utf-8'))
    assert multi['baseline_solution_sha256'] == digest(solution_path)
    stored_frontier = sorted(objective(plan) for plan in multi['frontier'])
    assert len(stored_frontier) == len(global_frontier)
    for actual, expected in zip(global_frontier, stored_frontier):
        assert_objective(actual, expected)
    multi_checks = [dict(plan_id=p['plan_id'], **verify_plan(p,boxes,drones,physics)) for p in multi['frontier']]
    endpoint_orderings = {'N_E_T':lambda x:x, 'N_T_E':lambda x:(x[0],x[2],x[1]),
                          'E_N_T':lambda x:(x[1],x[0],x[2]), 'T_N_E':lambda x:(x[2],x[0],x[1]),
                          'E_T_N':lambda x:(x[1],x[2],x[0]), 'T_E_N':lambda x:(x[2],x[1],x[0])}
    frontier_by_id = {plan['plan_id']:plan for plan in multi['frontier']}
    for key, ordering in endpoint_orderings.items():
        assert_objective(min(global_frontier,key=ordering), objective(frontier_by_id[multi['endpoints'][key]]))
    for representative in multi['representatives']:
        verify_plan(representative,boxes,drones,physics)
        assert_objective(objective(representative),objective(frontier_by_id[representative['plan_id']]))
    independent_frontier = dict(shared, method='Complete real-box-ID bitmask nondominated-label DP; each mask anchors its least-significant bit. Independently merged local frontiers.',
                                local_frontiers=local_frontiers, global_frontier=global_frontier, stage_statistics=statistics,
                                numeric_tolerances=dict(energy_kwh=ENERGY_TOL,time_s=TIME_TOL))
    write_json(multi_path.parent/'independent_frontier.json', independent_frontier)
    write_json(multi_path.parent/'independent_endpoints.json', dict(shared, method='Endpoints extracted from independently verified complete frontier.',
                endpoints={key:min(global_frontier,key=sortkey) for key,sortkey in endpoint_orderings.items()},
                six_scalar_endpoint_mappings_verified=True,endpoint_mapping_count=len(endpoint_orderings),
                representatives_match_frontier=True))
    write_json(multi_path.parent/'validation.json', dict(shared, formal_result_sha256=digest(multi_path),
                 complete_frontier_matches_independent_proof=True, frontier_point_count=len(global_frontier), plans=multi_checks))
    sensitivity_path = ROOT/'results/question1_reserve_sensitivity/sensitivity.json'
    sensitivity = json.loads(sensitivity_path.read_text(encoding='utf-8'))
    assert sensitivity['baseline_solution_sha256'] == digest(solution_path)
    quantities = QuantityAudit(boxes,drones,physics)
    capacity_path = sensitivity_path.parent/'all_capacity_feasible_batches.json'
    saved_pools = json.loads(capacity_path.read_text(encoding='utf-8'))
    critical_records = []
    single_limits = []
    for sid,pool in quantities.pools.items():
        saved = {(p['model'],tuple(p['counts'])):p for p in saved_pools[sid]}
        assert len(saved) == len(pool)
        for p in pool:
            old = saved[(p['model'],p['counts'])]
            assert abs(float(p['rho'])-old['critical_reserve_percent']) < 1e-9
            assert abs(float(p['energy'])-old['energy_kwh']) < 1e-9
            critical_records.append(dict(service_id=sid,model=p['model'],counts=p['counts'],critical_reserve_percent=float(p['rho']),critical_reserve_percent_40digit=str(p['rho'])))
        for k,count in enumerate(quantities.demands[sid]):
            if count:
                one = tuple(int(j == k) for j in range(len(TYPES)))
                best = max(p['rho'] for p in pool if p['counts'] == one)
                single_limits.append(dict(service_id=sid,type=TYPES[k],limit=best))
    ceiling = min(p['limit'] for p in single_limits)
    assert abs(float(ceiling)-sensitivity['global_feasibility_limit']['reserve_percent']) < 1e-9
    edges = sorted(set([Decimal(0),ceiling]+[p['rho'] for pool in quantities.pools.values() for p in pool if 0 < p['rho'] < ceiling]))
    assert len(edges)-2 == sensitivity['algorithm']['distinct_critical_values_within_feasible_range']
    event_checks = []
    for i,(left,right) in enumerate(zip(edges,edges[1:]),1):
        rho = (left+right)/2
        actual = quantities.global_objective(rho)
        plan = next(p for p in sensitivity['interval_plans'] if D(p['lower_percent']) <= rho <= D(p['upper_percent']))
        assert_objective(actual,objective(plan))
        event_checks.append(dict(lower_percent=float(left),upper_percent=float(right),checked_midpoint_percent=float(rho),plan_id=plan['plan_id'],objective=actual))
    endpoint_checks = []
    for rho in edges:
        actual = quantities.global_objective(rho)
        plan = next(p for p in sensitivity['interval_plans'] if float(rho) <= p['upper_percent']+1e-10)
        assert_objective(actual,objective(plan))
        endpoint_checks.append(dict(reserve_percent=float(rho),objective=actual,equality_feasible=True))
    assert quantities.global_objective(ceiling+D('0.000001')) is None
    scenario_checks = []
    for plan in sensitivity['scenarios']:
        value = quantities.global_objective(plan['reserve_percent'])
        assert (value is not None) == plan['feasible']
        item = dict(plan_id=plan['plan_id'],reserve_percent=plan['reserve_percent'],feasible=value is not None,independent_objective=value)
        if value is not None:
            assert_objective(value,objective(plan))
            item.update(verify_plan(plan,boxes,drones,physics))
        else:
            assert not plan['batches'] and all(plan[key] is None for key in ('flight_count','energy_kwh','work_time_s'))
        scenario_checks.append(item)
    interval_checks = [dict(plan_id=p['plan_id'],**verify_plan(p,boxes,drones,physics)) for p in sensitivity['interval_plans']]
    payload_checks = []
    maximum_payload_difference = 0.0
    samples = [(p['service_id'],p['model'],sample) for p in sensitivity['payloads'] for sample in p['samples']]
    samples.extend((p['service_id'],p['model'],p) for p in sensitivity['payload_curves'])
    for sid,model,sample in samples:
        value,label = physics.safe_payload(sid,model,sample['reserve_percent'])
        assert label == sample['limiting_factor']
        if value is None:
            assert sample['safe_payload_kg'] is None
            difference = None
        else:
            difference = abs(value-sample['safe_payload_kg'])
            maximum_payload_difference = max(maximum_payload_difference,difference)
            assert difference < 1.1e-6
        payload_checks.append(dict(service_id=sid,model=model,reserve_percent=sample['reserve_percent'],safe_payload_kg=value,limiting_factor=label,absolute_difference_kg=difference))
    for record in sensitivity['payloads']:
        drone = drones[record['model']]
        for field,mass in [('full_load_maximum_reserve_percent',drone['payload_kg']),('empty_roundtrip_maximum_reserve_percent',0)]:
            assert abs(float(physics.cost(record['service_id'],record['model'],D(mass),0)[-1])-record[field]) < 1e-9
    payload_report = dict(shared, formal_result_sha256=digest(sensitivity_path),method='Independent 40-digit Decimal energy and 1e-10 kg root brackets; all published samples and 0..60 percent integer curves.',
                           sample_record_count=len(samples),maximum_payload_difference_kg=maximum_payload_difference,payload_rows=payload_checks,
                           all_physical_batch_critical_reserves=critical_records)
    write_json(sensitivity_path.parent/'independent_payloads.json',payload_report)
    write_json(sensitivity_path.parent/'payload_validation.json',payload_report)
    write_json(sensitivity_path.parent/'independent_scenarios.json',dict(shared,method='Independent Decimal quantity DP after complete real-box-ID baseline/Pareto verification; every reported scenario.',
                 scenarios=scenario_checks,global_maximum_reserve_percent=float(ceiling),global_maximum_reserve_percent_40digit=str(ceiling)))
    final = dict(shared,formal_result_sha256=digest(sensitivity_path),method='Independent input reread, Decimal cost, candidate reconstruction, all critical intervals and equality endpoints; no main solver import.',
                 candidate_completeness=dict(physical_candidates=len(critical_records),all_candidates_matched=True),
                 event_scan=dict(elementary_interval_count=len(event_checks),endpoint_count=len(endpoint_checks),intervals=event_checks,endpoints=endpoint_checks),
                 global_feasibility_limit_percent_40digit=str(ceiling),global_limit_just_above_infeasible=True,
                 all_output_plan_checks=interval_checks+scenario_checks,payload_record_count=len(payload_checks),
                 maximum_payload_difference_kg=maximum_payload_difference,elapsed_seconds=time.perf_counter()-started)
    write_json(sensitivity_path.parent/'validation.json',final)
    print(json.dumps(dict(passed=True,baseline=baseline_check['recomputed_totals'],frontier=global_frontier,
                          physical_candidates=len(critical_records),elementary_intervals=len(event_checks),endpoints=len(endpoint_checks),
                          payload_records=len(payload_checks),global_limit_percent=float(ceiling),elapsed_seconds=final['elapsed_seconds']),ensure_ascii=False,indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--geometry-only',action='store_true')
    args = parser.parse_args()
    solution_path = ROOT/'results/question1_batching/solution.json'
    solution = json.loads(solution_path.read_text(encoding='utf-8'))
    paths,nodes,drones,boxes = read_originals(solution)
    routes,geometry = geometry_audit(solution_path,solution,paths,nodes)
    print('GEOMETRY',json.dumps(geometry['summary'],ensure_ascii=False),flush=True)
    if not args.geometry_only:
        full_audit(solution_path,solution,paths,drones,boxes,routes,geometry)


if __name__ == '__main__':
    main()
